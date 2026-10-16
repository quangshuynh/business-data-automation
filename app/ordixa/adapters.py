import csv
import difflib
import hashlib
import io
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from app.ordixa.issues import InputError
from app.ordixa.normalize import is_blank

# inputs are read into memory once so the recorded checksum always describes
# exactly the bytes that were parsed; larger exports should be split by period
MAX_INPUT_BYTES = 200 * 1024 * 1024


@dataclass(frozen=True)
class SourceRow:
    """one data row from an input file with its spreadsheet row number"""

    row: int
    values: dict


@dataclass(frozen=True)
class SourceTable:
    """rows read from one input file"""

    role: str
    path: Path
    sha256: str
    header: tuple[str, ...]
    rows: tuple[SourceRow, ...]
    malformed_rows: tuple[int, ...] = ()


def file_sha256(path):
    """
    calculates the SHA-256 fingerprint of a file
    :param path: file path
    :returns: hexadecimal digest
    """
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def match_inputs(input_dir, profile):
    """
    matches files to dataset roles and collects every discovery problem
    :param input_dir: directory containing customer exports
    :param profile: validated import profile
    :returns: tuple of role-to-paths mapping, ignored file names, and problems
    """
    input_dir = Path(input_dir)
    if not input_dir.is_dir():
        raise InputError(f"input directory not found: {input_dir}")

    claimed = {}
    discovered = {}
    problems = []
    for role, spec in profile.datasets.items():
        paths = set()
        for pattern in spec.files:
            for path in input_dir.glob(pattern):
                if path.is_file() and path.suffix.lower().lstrip(".") in (
                    profile.formats
                ):
                    paths.add(path)

        for path in sorted(paths, key=lambda path: path.name):
            if path.name in claimed:
                problems.append(
                    f"{path.name} matches the file patterns of both the "
                    f"{claimed[path.name]} and {role} datasets; make the profile "
                    "patterns more specific"
                )
                continue
            claimed[path.name] = role

        if not paths and spec.required:
            problems.append(
                f"no {role} file found in {input_dir} "
                f"(expected a file matching {list(spec.files)})"
            )
        discovered[role] = sorted(paths, key=lambda path: path.name)

    ignored = sorted(
        path.name
        for path in input_dir.iterdir()
        if path.is_file() and path.name not in claimed
    )
    return discovered, ignored, problems


def discover_inputs(input_dir, profile):
    """
    finds the files for each dataset role using the profile's file patterns
    :param input_dir: directory containing customer exports
    :param profile: validated import profile
    :returns: tuple of role-to-paths mapping and ignored file names
    """
    discovered, ignored, problems = match_inputs(input_dir, profile)
    if problems:
        raise InputError(problems[0])
    return discovered, ignored


def clean_header(cell):
    """
    normalizes a header cell
    :param cell: raw header value
    :returns: trimmed header text
    """
    if cell is None:
        return ""
    return str(cell).replace("﻿", "").strip()


def missing_columns(header, spec):
    """
    lists required mapped source columns absent from a header
    :param header: header names from the file
    :param spec: dataset specification
    :returns: list of quoted column names, each with the closest header found
    """
    missing = []
    for source in spec.source_columns():
        if source in header or source in spec.optional_columns:
            continue
        hint = difflib.get_close_matches(source, [h for h in header if h], n=1)
        missing.append(f"{source!r}" + (f" (found {hint[0]!r})" if hint else ""))
    return missing


def check_header(header, spec, path):
    """
    verifies that all required mapped source columns exist
    :param header: header names from the file
    :param spec: dataset specification
    :param path: input file path
    :returns: none
    """
    duplicates = sorted({name for name in header if name and header.count(name) > 1})
    if duplicates:
        raise InputError(f"{path.name} has duplicate column headers: {duplicates}")

    missing = missing_columns(header, spec)
    if missing:
        raise InputError(
            f"{path.name} ({spec.role}) is missing required columns: "
            f"{', '.join(missing)}"
        )


def read_csv_table(path, spec, encoding, data=None):
    """
    reads a CSV export keeping every cell as the exact source text
    :param path: CSV file path
    :param spec: dataset specification
    :param encoding: text encoding declared by the profile
    :param data: file bytes already read, or none to read the file
    :returns: tuple of rows, header, and malformed row numbers
    """
    if data is None:
        data = path.read_bytes()
    try:
        with io.StringIO(data.decode(encoding), newline="") as handle:
            reader = csv.reader(handle)
            header_cells = next(reader, None)
            if header_cells is None:
                raise InputError(f"{path.name} is empty")
            header = tuple(clean_header(cell) for cell in header_cells)
            check_header(header, spec, path)

            rows = []
            malformed = []
            for cells in reader:
                row_number = reader.line_num
                if not any(cell.strip() for cell in cells):
                    continue
                if len(cells) != len(header):
                    malformed.append(row_number)
                    continue
                rows.append(SourceRow(row_number, dict(zip(header, cells))))
    except UnicodeDecodeError as error:
        raise InputError(
            f"{path.name} is not valid {encoding} text: {error.reason}"
        ) from error
    except csv.Error as error:
        raise InputError(f"{path.name} could not be parsed as CSV: {error}") from error

    return rows, header, malformed


def read_xlsx_table(path, spec, data=None):
    """
    reads an XLSX export keeping typed cell values
    :param path: workbook path
    :param spec: dataset specification
    :param data: file bytes already read, or none to read the file
    :returns: tuple of rows, header, and malformed row numbers
    """
    if data is None:
        data = path.read_bytes()
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises many unrelated exception types
        raise InputError(f"{path.name} could not be opened as XLSX: {error}") from error

    try:
        if spec.sheet:
            if spec.sheet not in workbook.sheetnames:
                raise InputError(
                    f"{path.name} has no worksheet named {spec.sheet!r} "
                    f"(found {workbook.sheetnames})"
                )
            sheet = workbook[spec.sheet]
        else:
            sheet = workbook.worksheets[0]

        iterator = sheet.iter_rows(values_only=True)
        header_cells = next(iterator, None)
        if header_cells is None:
            raise InputError(f"{path.name} is empty")
        header = tuple(clean_header(cell) for cell in header_cells)
        while header and not header[-1]:
            header = header[:-1]
        check_header(header, spec, path)

        rows = []
        malformed = []
        for offset, cells in enumerate(iterator, start=2):
            cells = list(cells or ())
            if all(is_blank(cell) for cell in cells):
                continue
            if any(not is_blank(cell) for cell in cells[len(header) :]):
                malformed.append(offset)
                continue
            cells += [None] * (len(header) - len(cells))
            rows.append(SourceRow(offset, dict(zip(header, cells))))
    finally:
        workbook.close()

    return rows, header, malformed


def read_table(path, spec, profile):
    """
    reads one input file for a dataset role
    :param path: input file path
    :param spec: dataset specification
    :param profile: validated import profile
    :returns: source table
    """
    path = Path(path)
    try:
        size = path.stat().st_size
        if size > MAX_INPUT_BYTES:
            raise InputError(
                f"{path.name} is {size:,} bytes, above the "
                f"{MAX_INPUT_BYTES // (1024 * 1024)} MB input limit; split the "
                "export by reporting period"
            )
        data = path.read_bytes()
    except OSError as error:
        raise InputError(f"{path.name} could not be read: {error.strerror}") from error

    sha256 = hashlib.sha256(data).hexdigest()
    if path.suffix.lower() == ".xlsx":
        rows, header, malformed = read_xlsx_table(path, spec, data)
    else:
        rows, header, malformed = read_csv_table(path, spec, profile.csv_encoding, data)

    return SourceTable(
        role=spec.role,
        path=path,
        sha256=sha256,
        header=header,
        rows=tuple(rows),
        malformed_rows=tuple(malformed),
    )
