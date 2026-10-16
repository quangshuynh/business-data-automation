import math
import re
import textwrap
from datetime import date, datetime
from decimal import Decimal

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.ordixa.issues import OUTCOME_LABELS
from app.ordixa.normalize import identifier_sort_key

SHEET_NAMES = ("Summary", "Exceptions", "Rejected Rows", "Full Ledger")

# Ordixa palette (docs/BRAND.md): graphite headers with off-white text and a
# chartreuse rule, teal-graphite labels on white, tinted status chips; every
# text/background pair below has at least a 7:1 contrast ratio (WCAG AAA)
GRAPHITE = "111D22"
SURFACE = "22343A"
CHARTREUSE = "D9F17A"
OFF_WHITE = "F2F4EF"
AMBER = "E6AE65"
MUTED = "3D4F55"
RULE = "C9D1CC"
FONT_NAME = "Calibri"
MONO_FONT_NAME = "Consolas"
STATUS_STYLES = {
    "paid": ("Paid", "E2EFE6", "1E4D35"),
    "partial": ("Partially paid", "F7E6C8", "5E3A0B"),
    "unpaid": ("Unpaid", "F5DAD5", "7A2216"),
    "overpaid": ("Overpaid", "DCEAED", "1B4855"),
}
SEVERITY_STYLES = {
    "error": ("F5DAD5", "7A2216"),
    "warning": ("F7E6C8", "5E3A0B"),
    "excluded": ("E9ECE8", SURFACE),
}
TAB_COLORS = dict(zip(SHEET_NAMES, (GRAPHITE, AMBER, "B5523F", SURFACE)))

MONEY_FORMAT = "#,##0.00;-#,##0.00"
DATE_FORMAT = "yyyy-mm-dd"
INTEGER_FORMAT = "#,##0"

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n", "＝", "＋")
# Excel stores numbers as binary doubles and openpyxl writes them through float
# repr, so only values with at most 15 significant digits survive exactly
EXCEL_SAFE_DIGITS = 15
MAX_CELL_TEXT = 32_767
TRUNCATION_MARK = " [truncated]"
# identifiers are always text cells so leading zeros and long numbers survive
TEXT_COLUMNS = {"order_id", "customer_id", "record_id", "order_ref"}
PLAIN_NUMBER = re.compile(r"^[-+]?[$€£¥]?\s?[0-9][0-9.,\s]*$")

HEADER_FONT = Font(name=FONT_NAME, bold=True, color=OFF_WHITE)
HEADER_FILL = PatternFill("solid", fgColor=GRAPHITE)
HEADER_RULE = Border(bottom=Side(style="medium", color=CHARTREUSE))
BRAND_FONT = Font(name=FONT_NAME, bold=True, size=20, color=CHARTREUSE)
TITLE_FONT = Font(name=FONT_NAME, size=13, color=OFF_WHITE)
SECTION_FONT = Font(name=FONT_NAME, bold=True, size=12, color=GRAPHITE)
LABEL_FONT = Font(name=FONT_NAME, color=MUTED)
VALUE_FONT = Font(name=FONT_NAME, color=GRAPHITE)
BOLD_VALUE_FONT = Font(name=FONT_NAME, bold=True, color=GRAPHITE)
HASH_FONT = Font(name=MONO_FONT_NAME, size=9, color=MUTED)
BOTTOM_RULE = Border(bottom=Side(style="thin", color=RULE))
WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")
# figures and dates sit one character in from the right edge of their column
FIGURE = Alignment(horizontal="right", vertical="top", indent=1)
# long free-text columns wrap at these widths instead of stretching the sheet
WRAPPED_COLUMNS = {
    "review_note": 46,
    "message": 58,
    "discrepancy_flags": 34,
    "transaction_refs": 28,
    "original_values": 60,
}
LINE_HEIGHT = 15

LEDGER_HEADERS = {
    "order_ref": "Order",
    "order_id": "Order ID",
    "order_date": "Order date",
    "customer_id": "Customer ID",
    "customer_name": "Customer",
    "customer_email": "Customer email",
    "customer_found": "Customer on file",
    "order_total": "Order total",
    "payments_received": "Payments",
    "refunds": "Refunds",
    "adjustments": "Adjustments",
    "net_paid": "Net paid",
    "balance_due": "Balance due",
    "outstanding_balance": "Outstanding",
    "overpayment_amount": "Overpaid",
    "financial_status": "Status",
    "discrepancy_flags": "Discrepancy flags",
    "transaction_count": "Transactions",
    "transaction_refs": "Transaction IDs",
    "last_payment_date": "Last payment date",
    "source_location": "Source",
}
EXCEPTION_COLUMNS = [
    "exception_type",
    "order_ref",
    "order_date",
    "customer_name",
    "financial_status",
    "order_total",
    "net_paid",
    "outstanding_balance",
    "overpayment_amount",
    "review_note",
    "transaction_count",
    "transaction_refs",
    "discrepancy_flags",
    "source_location",
]
EXCEPTION_HEADERS = {
    **LEDGER_HEADERS,
    "exception_type": "Exception",
    "review_note": "What to check",
}
EXCEPTION_ORDER = {
    "Unpaid": 0,
    "Partially paid": 1,
    "Overpaid": 2,
    "Paid with discrepancy": 3,
}
MONEY_COLUMNS = {
    "order_total",
    "payments_received",
    "refunds",
    "adjustments",
    "net_paid",
    "balance_due",
    "outstanding_balance",
    "overpayment_amount",
}
# right-aligned headers so they sit above their figures
NUMERIC_COLUMNS = MONEY_COLUMNS | {
    "transaction_count",
    "row",
    "order_date",
    "last_payment_date",
}
# plain-language guidance shown beside each exception on the Exceptions sheet
REVIEW_NOTES = {
    "Unpaid": (
        "No counted payment. Confirm whether the order was invoiced, cancelled, "
        "or paid outside these exports."
    ),
    "Partially paid": (
        "Net received is below the order total. Collect the balance, or confirm "
        "the refund or discount that explains it."
    ),
    "Overpaid": (
        "More was received than the order total. Check for a duplicate charge, "
        "then refund or credit the difference."
    ),
    "Paid with discrepancy": (
        "Amounts match the order total, but a discrepancy flag needs a look."
    ),
}


def safe_text(value):
    """
    neutralizes untrusted text so spreadsheet tools never evaluate it

    text that could start a formula gets a leading apostrophe, except values
    that are plainly numbers such as -$89.00; XML-illegal control characters
    are removed
    :param value: untrusted text
    :returns: sanitized text
    """
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    if text.startswith(FORMULA_PREFIXES) and not PLAIN_NUMBER.match(text):
        return "'" + text
    return text


def exact_in_excel(value):
    """
    checks whether a number can be stored as an Excel number without losing
    digits
    :param value: decimal or integer
    :returns: true when at most 15 significant digits are needed
    """
    if isinstance(value, Decimal):
        return value.is_finite() and len(value.as_tuple().digits) <= EXCEL_SAFE_DIGITS
    return abs(value) < 10**EXCEL_SAFE_DIGITS


def imprecise_amounts(result):
    """
    counts monetary values too precise for an Excel number cell
    :param result: run result
    :returns: number of summary and ledger values written as exact text
    """
    values = list(result.summary.values())
    for row in result.ledger.to_dict(orient="records"):
        values.extend(row[column] for column in MONEY_COLUMNS)
    return sum(
        1
        for value in values
        if isinstance(value, Decimal) and not exact_in_excel(value)
    )


def cell_text(value):
    """
    sanitizes text for a cell and keeps it within Excel's length limit
    :param value: untrusted text
    :returns: text no longer than 32,767 characters
    """
    text = safe_text(value)
    if len(text) > MAX_CELL_TEXT:
        text = text[: MAX_CELL_TEXT - len(TRUNCATION_MARK)] + TRUNCATION_MARK
    return text


def status_label(status):
    """
    converts a canonical financial status into its display label
    :param status: canonical status
    :returns: display label
    """
    return STATUS_STYLES.get(status, (str(status).title(), None, None))[0]


def write_cell(sheet, row, column, value, column_name=None):
    """
    writes one value with a type-appropriate number format
    :param sheet: worksheet
    :param row: row number
    :param column: column number
    :param value: python value
    :param column_name: canonical column name used to detect money columns
    :returns: written cell
    """
    cell = sheet.cell(row=row, column=column)

    if value is None or (not isinstance(value, str) and pd.isna(value)):
        cell.value = None
    elif isinstance(value, bool):
        cell.value = "Yes" if value else "No"
    elif column_name in TEXT_COLUMNS:
        cell.value = cell_text(value)
        cell.data_type = "s"
    elif isinstance(value, (Decimal, int)) and not exact_in_excel(value):
        # an Excel number would silently drop digits; keep the exact value as
        # right-aligned text instead (the Summary sheet counts these)
        cell.value = format(value, "f") if isinstance(value, Decimal) else str(value)
        cell.data_type = "s"
        cell.alignment = Alignment(horizontal="right")
    elif isinstance(value, Decimal):
        # at most 15 significant digits, so the float written by openpyxl
        # converts back to exactly this decimal
        cell.value = value
        cell.number_format = MONEY_FORMAT
    elif isinstance(value, int):
        cell.value = value
        cell.number_format = MONEY_FORMAT if column_name in MONEY_COLUMNS else "0"
    elif isinstance(value, datetime):
        cell.value = value.replace(tzinfo=None)
        cell.number_format = "yyyy-mm-dd hh:mm:ss"
    elif isinstance(value, date):
        cell.value = value
        cell.number_format = DATE_FORMAT
    else:
        cell.value = cell_text(value)
        cell.data_type = "s"

    return cell


def style_header(cell, numeric=False):
    """
    applies the graphite header style with its chartreuse rule
    :param cell: header cell
    :param numeric: right-align the header above a column of figures
    :returns: none
    """
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.border = HEADER_RULE
    cell.alignment = Alignment(
        horizontal="right" if numeric else "left",
        vertical="center",
        wrap_text=True,
        indent=1 if numeric else 0,
    )


def write_table(sheet, start_row, columns, headers, rows):
    """
    writes a header row and records as a filterable table
    :param sheet: worksheet
    :param start_row: row of the header
    :param columns: canonical column names
    :param headers: display headers in column order
    :param rows: list of dictionaries keyed by canonical column
    :returns: last row written
    """
    for index, (column, header) in enumerate(zip(columns, headers), start=1):
        cell = sheet.cell(row=start_row, column=index, value=header)
        style_header(cell, numeric=column in NUMERIC_COLUMNS)
    sheet.row_dimensions[start_row].height = 30

    row_number = start_row
    for record in rows:
        row_number += 1
        for index, column in enumerate(columns, start=1):
            cell = write_cell(sheet, row_number, index, record.get(column), column)
            if column in WRAPPED_COLUMNS:
                cell.alignment = WRAP
            elif cell.alignment.horizontal == "right" or (
                isinstance(cell.value, (int, float, Decimal, date))
                and not isinstance(cell.value, bool)
            ):
                cell.alignment = FIGURE
            else:
                cell.alignment = TOP
            if column == "financial_status" and record.get(column) in STATUS_STYLES:
                label, fill, color = STATUS_STYLES[record[column]]
                cell.value = label
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.font = Font(name=FONT_NAME, color=color, bold=True)

    last_column = get_column_letter(len(columns))
    sheet.auto_filter.ref = f"A{start_row}:{last_column}{max(row_number, start_row)}"
    sheet.freeze_panes = sheet.cell(
        row=start_row + 1, column=2 if len(columns) > 3 else 1
    )
    return row_number


def fit_columns(sheet, columns=(), minimum=9, maximum=60):
    """
    sizes columns to their content within readable limits; long free-text
    columns get a fixed width and wrap instead
    :param sheet: worksheet
    :param columns: canonical column names in sheet order, for wrapped widths
    :param minimum: smallest width
    :param maximum: largest width
    :returns: none
    """
    widths = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            if isinstance(cell.value, Decimal):
                length = len(f"{cell.value:,.2f}") + 1
            elif isinstance(cell.value, (date, datetime)):
                length = 11
            else:
                length = max(len(line) for line in str(cell.value).splitlines() or [""])
            if cell.font is not None and cell.font.b:
                # bold text is wider than regular body text
                length = int(length * 1.1) + 1
            widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), length)

    for letter, width in widths.items():
        sheet.column_dimensions[letter].width = max(minimum, min(maximum, width + 2))
    for index, column in enumerate(columns, start=1):
        if column in WRAPPED_COLUMNS:
            letter = get_column_letter(index)
            sheet.column_dimensions[letter].width = min(
                WRAPPED_COLUMNS[column], max(minimum, widths.get(letter, 0) + 2)
            )


def fit_rows(sheet, columns, start_row=2):
    """
    sets row heights for wrapped text, which spreadsheet applications do not
    reliably grow on their own
    :param sheet: worksheet whose column widths are already set
    :param columns: canonical column names in sheet order
    :param start_row: first data row
    :returns: none
    """
    wrapped = [
        (index, sheet.column_dimensions[get_column_letter(index)].width)
        for index, column in enumerate(columns, start=1)
        if column in WRAPPED_COLUMNS
    ]
    for row_number in range(start_row, sheet.max_row + 1):
        lines = 1
        for index, width in wrapped:
            value = sheet.cell(row=row_number, column=index).value
            if isinstance(value, str) and len(value) > width - 2:
                # body text averages slightly narrower than the width unit
                chars = max(1, math.floor(width * 1.1) - 2)
                lines = max(lines, len(textwrap.wrap(value, chars)) or 1)
        if lines > 1:
            sheet.row_dimensions[row_number].height = LINE_HEIGHT * lines + 2


def set_print_layout(sheet, landscape=True, title_rows=None):
    """
    prepares a sheet for printing: one page wide, repeated headers, footer
    :param sheet: worksheet
    :param landscape: print in landscape orientation
    :param title_rows: rows repeated on every printed page, such as "1:1"
    :returns: none
    """
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = sheet.page_margins.right = 0.4
    sheet.page_margins.top = sheet.page_margins.bottom = 0.5
    if title_rows:
        sheet.print_title_rows = title_rows
    sheet.oddFooter.left.text = "Ordixa"
    sheet.oddFooter.left.size = 8
    sheet.oddFooter.right.text = "&A - page &P of &N"
    sheet.oddFooter.right.size = 8
    sheet.sheet_properties.tabColor = TAB_COLORS[sheet.title]


def exception_type(row):
    """
    classifies a ledger row for the exceptions sheet
    :param row: ledger row dictionary
    :returns: exception label
    """
    if row["financial_status"] == "paid":
        return "Paid with discrepancy"
    return status_label(row["financial_status"])


def build_summary_sheet(sheet, result, metadata):
    """
    writes headline totals, quality counts, guidance, and processing metadata
    :param sheet: worksheet
    :param result: run result
    :param metadata: execution metadata for the run
    :returns: none
    """
    summary = result.summary
    profile = result.profile
    currency = profile.amounts.currency
    width = 6

    for brand_row in (1, 2):
        for column in range(1, width + 1):
            sheet.cell(row=brand_row, column=column).fill = HEADER_FILL
    for column in range(1, width + 1):
        sheet.cell(row=2, column=column).border = HEADER_RULE
    sheet["A1"] = "Ordixa"
    sheet["A1"].font = BRAND_FONT
    sheet["A1"].alignment = Alignment(vertical="bottom")
    sheet["A2"] = "Reconciliation summary"
    sheet["A2"].font = TITLE_FONT
    sheet["A2"].alignment = Alignment(vertical="top")
    sheet.row_dimensions[1].height = 32
    sheet.row_dimensions[2].height = 22
    sheet["A3"] = safe_text(
        f"Profile {profile.name} v{profile.version} - amounts in {currency} - "
        f"generated {metadata['generated_at']} UTC"
    )
    sheet["A3"].font = LABEL_FONT
    sheet.row_dimensions[3].height = 20
    row = 5

    def section(title):
        nonlocal row
        sheet.cell(row=row, column=1, value=title).font = SECTION_FONT
        for column in range(1, width + 1):
            sheet.cell(row=row, column=column).border = BOTTOM_RULE
        sheet.row_dimensions[row].height = 20
        row += 1

    def line(label, value, number_format=None, emphasis=False):
        nonlocal row
        sheet.cell(row=row, column=1, value=label).font = LABEL_FONT
        cell = write_cell(sheet, row, 2, value)
        if number_format:
            cell.number_format = number_format
        cell.font = BOLD_VALUE_FONT if emphasis else VALUE_FONT
        # figures align right; text runs left into the empty columns beside it
        if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
            cell.alignment = Alignment(horizontal="right")
        row += 1

    def note(label, text):
        nonlocal row
        sheet.cell(row=row, column=1, value=label).font = LABEL_FONT
        sheet.cell(row=row, column=2, value=text).font = VALUE_FONT
        row += 1

    def table_header(headers, numeric_from):
        nonlocal row
        for index, header in enumerate(headers, start=1):
            style_header(
                sheet.cell(row=row, column=index, value=header),
                numeric=index >= numeric_from,
            )
        row += 1

    section("Financial totals")
    line("Orders reconciled", summary["orders"], INTEGER_FORMAT)
    line("Order value", summary["order_value"], emphasis=True)
    line("Transactions counted", summary["transactions"], INTEGER_FORMAT)
    line("Payments received", summary["payments_received"])
    line("Refunds (deducted)", summary["refunds"])
    line("Adjustments (signed)", summary["adjustments"])
    line("Net received", summary["net_received"], emphasis=True)
    line("Reconciled to orders", summary["applied_to_orders"])
    line("Outstanding balances", summary["outstanding_balance"], emphasis=True)
    line("Overpayments", summary["overpayments"], emphasis=True)
    row += 1

    section("Exceptions")
    for status in ("unpaid", "partial", "overpaid"):
        line(f"{status_label(status)} orders", summary["status_counts"][status])
        _, fill, color = STATUS_STYLES[status]
        label_cell = sheet.cell(row=row - 1, column=1)
        label_cell.fill = PatternFill("solid", fgColor=fill)
        label_cell.font = Font(name=FONT_NAME, color=color, bold=True)
    line("Orders with discrepancy flags", summary["flagged_orders"])
    line("Orders needing review", summary["exception_orders"], emphasis=True)
    line("Paid orders", summary["status_counts"]["paid"])
    row += 1

    section("Reading this workbook")
    note("Exceptions sheet", "Orders to follow up, each with a note on what to check.")
    note(
        "Rejected Rows sheet", "Every row not used as-is, with its source row and why."
    )
    note("Full Ledger sheet", "Every reconciled order with its payments and refunds.")
    note("Rejected", "Not used in any total until the source export is fixed.")
    note(
        "Kept with warning", "Used in the totals; a detail such as contact data is off."
    )
    note(
        "Excluded by profile", "Deliberately ignored, such as a declined card payment."
    )
    note("Calculations", "Exact decimal arithmetic by Ordixa; no spreadsheet formulas.")
    row += 1

    section("Data quality")
    table_header(
        ["Dataset", "Rows read", "Accepted", "With warnings", "Rejected", "Excluded"],
        numeric_from=2,
    )
    for dataset, counts in result.counts.items():
        values = [
            dataset.title(),
            counts["rows"],
            counts["accepted"],
            counts["accepted_with_warnings"],
            counts["rejected"],
            counts["excluded"],
        ]
        for index, value in enumerate(values, start=1):
            write_cell(sheet, row, index, value).font = VALUE_FONT
        row += 1
    row += 1

    section("Processing")
    line("Run ID", metadata["run_id"])
    line("Generated (UTC)", metadata["generated_at"])
    line("Profile", f"{profile.name} {profile.version}")
    line("Application version", metadata["code_version"])
    line("Status", metadata["status"])
    line("Rejected rows (blocking)", summary["rejected_rows"], INTEGER_FORMAT)
    line("Rows kept with warnings", summary["warning_rows"], INTEGER_FORMAT)
    line("Rows excluded by profile", summary["excluded_rows"], INTEGER_FORMAT)
    imprecise = imprecise_amounts(result)
    if imprecise:
        line("Amounts stored as exact text (over 15 digits)", imprecise, INTEGER_FORMAT)
    row += 1

    section("Input files")
    table_header(["Role", "File", "Rows", "SHA-256", "", ""], numeric_from=3)
    sheet.cell(row=row - 1, column=4).alignment = Alignment(horizontal="left", indent=1)
    for input_file in result.inputs:
        for index, value in enumerate(
            [input_file.role, input_file.name, input_file.data_rows, input_file.sha256],
            start=1,
        ):
            cell = write_cell(sheet, row, index, value)
            cell.font = HASH_FONT if index == 4 else VALUE_FONT
            if index == 3:
                cell.alignment = FIGURE
            elif index == 4:
                cell.alignment = Alignment(indent=1)
        row += 1

    for letter, column_width in zip("ABCDEF", (34, 24, 13, 15, 13, 13)):
        sheet.column_dimensions[letter].width = column_width
    sheet.sheet_view.showGridLines = False
    set_print_layout(sheet, landscape=False)


def build_exceptions_sheet(sheet, result):
    """
    writes orders that need follow-up with a note on what to check
    :param sheet: worksheet
    :param result: run result
    :returns: none
    """
    rows = []
    for row in result.exceptions.to_dict(orient="records"):
        kind = exception_type(row)
        rows.append({**row, "exception_type": kind, "review_note": REVIEW_NOTES[kind]})
    rows.sort(
        key=lambda item: (
            EXCEPTION_ORDER[item["exception_type"]],
            identifier_sort_key(item["order_id"]),
        )
    )

    headers = [EXCEPTION_HEADERS[column] for column in EXCEPTION_COLUMNS]
    write_table(sheet, 1, EXCEPTION_COLUMNS, headers, rows)
    for row_number, record in enumerate(rows, start=2):
        if record["financial_status"] in STATUS_STYLES:
            _, _, color = STATUS_STYLES[record["financial_status"]]
            sheet.cell(row=row_number, column=1).font = Font(
                name=FONT_NAME, color=color, bold=True
            )
    fit_columns(sheet, EXCEPTION_COLUMNS)
    fit_rows(sheet, EXCEPTION_COLUMNS)
    set_print_layout(sheet, title_rows="1:1")


def build_rejected_rows_sheet(sheet, result):
    """
    writes every row-level issue with its source location and outcome
    :param sheet: worksheet
    :param result: run result
    :returns: none
    """
    columns = [
        "dataset",
        "file",
        "row",
        "record_id",
        "outcome",
        "blocking",
        "code",
        "field",
        "value",
        "message",
        "original_values",
    ]
    headers = [
        "Dataset",
        "Source file",
        "Source row",
        "Record ID",
        "Outcome",
        "Prevented processing",
        "Issue code",
        "Field",
        "Value",
        "Reason",
        "Original values",
    ]
    rows = [
        {
            "dataset": issue.dataset,
            "file": issue.file,
            "row": issue.row,
            "record_id": issue.record_id,
            "outcome": OUTCOME_LABELS[issue.severity],
            "blocking": issue.blocking,
            "code": issue.code,
            "field": issue.field,
            "value": issue.value,
            "message": issue.message,
            "original_values": issue.original_values,
        }
        for issue in result.issues
    ]
    last_row = write_table(sheet, 1, columns, headers, rows)

    for row_number, issue in zip(range(2, last_row + 1), result.issues):
        fill, color = SEVERITY_STYLES[issue.severity]
        cell = sheet.cell(row=row_number, column=5)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name=FONT_NAME, color=color, bold=True)

    fit_columns(sheet, columns, maximum=70)
    fit_rows(sheet, columns)
    set_print_layout(sheet, title_rows="1:1")


def build_ledger_sheet(sheet, result):
    """
    writes the full reconciled ledger
    :param sheet: worksheet
    :param result: run result
    :returns: none
    """
    columns = list(LEDGER_HEADERS)
    rows = result.ledger.to_dict(orient="records")
    write_table(sheet, 1, columns, [LEDGER_HEADERS[c] for c in columns], rows)
    fit_columns(sheet, columns)
    fit_rows(sheet, columns)
    set_print_layout(sheet, title_rows="1:1")


def write_workbook(result, path, metadata):
    """
    writes the exceptions workbook for a run
    :param result: run result
    :param path: destination .xlsx path
    :param metadata: execution metadata (run_id, generated_at, code_version, status)
    :returns: destination path
    """
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = SHEET_NAMES[0]
    build_summary_sheet(summary_sheet, result, metadata)
    build_exceptions_sheet(workbook.create_sheet(SHEET_NAMES[1]), result)
    build_rejected_rows_sheet(workbook.create_sheet(SHEET_NAMES[2]), result)
    build_ledger_sheet(workbook.create_sheet(SHEET_NAMES[3]), result)

    workbook.properties.creator = "Ordixa"
    workbook.properties.title = safe_text(f"Ordixa run {metadata['run_id']}")
    workbook.save(path)
    return path
