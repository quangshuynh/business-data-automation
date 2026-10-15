"""
read-only inspection of unfamiliar CSV/XLSX exports before a profile exists

everything inferred here is a suggestion for the person writing a profile; it
is never used to process data, and no cell values are printed unless a
redacted shape preview is explicitly requested
"""

import codecs
import csv
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from app.ordixa.normalize import TEXT_ID_PATTERN, is_blank

SUPPORTED_SUFFIXES = {".csv": "csv", ".xlsx": "xlsx"}
UNSUPPORTED_SUFFIXES = {
    ".xls": "legacy Excel 97-2003 workbooks are not supported; re-save as .xlsx",
    ".xlsm": "macro-enabled workbooks are not supported; save a copy as .xlsx",
    ".xlsb": "binary Excel workbooks are not supported; re-save as .xlsx",
    ".ods": "OpenDocument spreadsheets are not supported; re-save as .xlsx",
    ".numbers": "Numbers documents are not supported; export as .xlsx or .csv",
    ".tsv": "tab-separated files are not supported; export as comma-separated CSV",
    ".txt": "text exports are not supported; export as comma-separated CSV",
}
DEFAULT_MAX_ROWS = 50_000
MAX_FILE_BYTES = 200 * 1024 * 1024
MAX_PREVIEW_ROWS = 10
MAX_DISTINCT_TRACKED = 50_000
EXCEL_DIGITS = 15

# candidate formats tried only to describe a column, never to parse data
CANDIDATE_DATE_FORMATS = (
    "%Y-%m-%d",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M:%S %z",
    "%Y-%m-%dT%H:%M:%S",
    "%Y-%m-%dT%H:%M:%S%z",
    "%Y/%m/%d",
    "%m/%d/%Y",
    "%d/%m/%Y",
    "%m/%d/%Y %H:%M",
    "%d/%m/%Y %H:%M",
    "%d.%m.%Y",
    "%b %d, %Y",
    "%d %b %Y",
)
DAY_MONTH_PAIRS = (
    ("%m/%d/%Y", "%d/%m/%Y"),
    ("%m/%d/%Y %H:%M", "%d/%m/%Y %H:%M"),
)
CURRENCY_SYMBOLS = "$€£¥"
AMOUNT_US = re.compile(r"^(?:[0-9]{1,3}(?:,[0-9]{3})+|[0-9]+)(?:\.(?P<f>[0-9]+))?$")
AMOUNT_EU = re.compile(r"^(?:[0-9]{1,3}(?:\.[0-9]{3})+|[0-9]+),(?P<f>[0-9]+)$")
DIGITS = re.compile(r"^[0-9]+$")
# cheap pre-filter so only date-shaped text is tried against every format
DATE_SHAPED = re.compile(
    r"^(?:[0-9]{1,4}[-/.][0-9]{1,2}[-/.][0-9]{2,4}|[0-9]{1,2} [A-Za-z]{3}|"
    r"[A-Za-z]{3} [0-9]{1,2},)"
)
EMAIL_LIKE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
PHONE_LIKE = re.compile(r"^\+?[0-9 ().-]{7,20}$")
ID_HEADER = re.compile(
    r"(?:^|[\s_-])(?:id|number|no|num|ref|reference|code)\.?$|#$", re.IGNORECASE
)
AMOUNT_HEADER = re.compile(
    r"amount|total|price|gross|net|paid|subtotal|tax|fee|refund|balance",
    re.IGNORECASE,
)


@dataclass
class ColumnStats:
    """running observations for one column"""

    name: str
    position: int
    values: int = 0
    blanks: int = 0
    kinds: Counter = field(default_factory=Counter)
    distinct: set = field(default_factory=set)
    distinct_overflow: bool = False
    date_hits: Counter = field(default_factory=Counter)
    date_values: int = 0
    amount_styles: Counter = field(default_factory=Counter)
    max_decimals: int = 0
    symbols: set = field(default_factory=set)
    negatives: set = field(default_factory=set)
    leading_zeros: int = 0
    padded: int = 0
    long_numbers: int = 0
    max_length: int = 0

    def observe(self, value):
        """
        records one cell
        :param value: raw cell value
        :returns: none
        """
        if is_blank(value):
            self.blanks += 1
            return

        self.values += 1
        if not self.distinct_overflow:
            self.distinct.add(value if isinstance(value, str) else repr(value))
            if len(self.distinct) > MAX_DISTINCT_TRACKED:
                self.distinct_overflow = True
                self.distinct.clear()

        if isinstance(value, bool):
            self.kinds["boolean"] += 1
        elif isinstance(value, datetime):
            self.kinds["date cell"] += 1
        elif isinstance(value, date):
            self.kinds["date cell"] += 1
        elif isinstance(value, (int, float, Decimal)):
            self.kinds["number cell"] += 1
            digits = len(str(value).replace("-", "").replace(".", "").lstrip("0"))
            if digits > EXCEL_DIGITS:
                self.long_numbers += 1
        else:
            self.observe_text(str(value))

    def observe_text(self, raw):
        """
        classifies one text cell
        :param raw: cell text
        :returns: none
        """
        text = raw.strip()
        if text != raw:
            self.padded += 1
        self.max_length = max(self.max_length, len(text))

        matching = set()
        if len(text) <= 40 and DATE_SHAPED.match(text):
            for date_format in CANDIDATE_DATE_FORMATS:
                try:
                    datetime.strptime(text, date_format)
                except ValueError:
                    continue
                matching.add(date_format)
        if matching:
            self.date_values += 1
            self.date_hits.update(matching)

        kind = "text"
        if DIGITS.match(text):
            kind = "digits"
            if len(text) > 1 and text.startswith("0"):
                self.leading_zeros += 1
            if len(text.lstrip("0")) > EXCEL_DIGITS:
                self.long_numbers += 1
        elif self.observe_amount(text):
            kind = "amount"
        elif matching:
            kind = "date text"
        elif EMAIL_LIKE.match(text):
            kind = "email"
        elif PHONE_LIKE.match(text) and sum(c.isdigit() for c in text) >= 7:
            kind = "phone"
        elif TEXT_ID_PATTERN.match(text) and any(c.isdigit() for c in text):
            kind = "identifier"
        self.kinds[kind] += 1

    def observe_amount(self, text):
        """
        checks whether text looks like a monetary amount and notes its style
        :param text: trimmed cell text
        :returns: true when the text is amount-shaped
        """
        body = text
        negative = None
        if body.startswith("(") and body.endswith(")"):
            negative, body = "parentheses", body[1:-1].strip()
        if body.startswith("-"):
            negative, body = "minus", body[1:].strip()
        symbol = ""
        if body[:1] and body[:1] in CURRENCY_SYMBOLS:
            symbol, body = body[0], body[1:].strip()
        elif body[-1:] and body[-1:] in CURRENCY_SYMBOLS:
            symbol, body = body[-1], body[:-1].strip()
        if body.startswith("-") and negative is None:
            negative, body = "minus", body[1:].strip()

        for style, pattern in (("1,234.56", AMOUNT_US), ("1.234,56", AMOUNT_EU)):
            match = pattern.match(body)
            if match:
                if not (symbol or negative or match.group("f") or "," in body):
                    return False
                self.amount_styles[style] += 1
                self.max_decimals = max(self.max_decimals, len(match.group("f") or ""))
                if symbol:
                    self.symbols.add(symbol)
                if negative:
                    self.negatives.add(negative)
                return True
        return False

    def summary(self, rows):
        """
        produces the reported description and suggestions
        :param rows: number of data rows scanned
        :returns: dictionary
        """
        kinds = {kind: count for kind, count in sorted(self.kinds.items())}
        distinct = None if self.distinct_overflow else len(self.distinct)
        suggestions = []
        notes = []
        values = self.values
        numeric_kinds = self.kinds["digits"] + self.kinds["number cell"]
        id_shaped = self.kinds["identifier"] + numeric_kinds

        date_like = self.kinds["date cell"] + self.date_values
        amount_like = self.kinds["amount"] + numeric_kinds
        common_kind, common_count = (
            self.kinds.most_common(1)[0] if self.kinds else ("", 0)
        )
        if values == 0:
            basic = "empty"
        elif date_like == values:
            basic = "date"
        elif self.kinds["number cell"] == values:
            basic = "number"
        elif self.kinds["digits"] == values:
            basic = "digits"
        elif self.kinds["amount"] and amount_like == values:
            basic = "amount"
        elif self.kinds["identifier"] and id_shaped == values:
            basic = "identifier"
        elif date_like * 2 >= values:
            basic = "date"
            notes.append(
                f"{values - date_like} value(s) match no common date format, "
                "for example an impossible or malformed date"
            )
        elif self.kinds["amount"] and amount_like * 2 >= values:
            basic = "amount"
            notes.append(f"{values - amount_like} value(s) are not amounts")
        elif common_count == values:
            basic = common_kind
        elif common_count * 5 >= values * 4:
            basic = common_kind
            notes.append(f"{values - common_count} value(s) are not {common_kind}")
        else:
            basic = "mixed"
            notes.append("values of several kinds; check for stray totals or notes")

        if (
            values
            and id_shaped == values
            and (
                ID_HEADER.search(self.name)
                or basic in {"digits", "identifier", "mixed"}
            )
        ):
            uniqueness = (
                "unique in scanned rows"
                if distinct == values
                else f"{values - distinct} repeated value(s): a reference column, "
                "or a key with duplicate rows"
                if distinct is not None
                else "too many values to check uniqueness"
            )
            suggestions.append(f"possible identifier ({uniqueness})")
        if basic == "date":
            if self.kinds["date cell"]:
                suggestions.append(
                    f"possible date ({self.kinds['date cell']} native date cell(s))"
                )
            if self.date_values:
                complete = sorted(
                    name
                    for name, count in self.date_hits.items()
                    if count == self.date_values
                )
                partial = ", ".join(
                    f"{name} ({count} of {self.date_values})"
                    for name, count in sorted(self.date_hits.items())
                )
                suggestions.append(
                    "possible date; formats matching every date-shaped value: "
                    + ", ".join(complete)
                    if complete
                    else "possible date; no single format fits every value, so a "
                    "profile would need several: " + partial
                )
                for us, eu in DAY_MONTH_PAIRS:
                    if us in complete and eu in complete:
                        notes.append(
                            f"day/month order is ambiguous ({us} and {eu} both "
                            "fit); confirm with the business before choosing"
                        )
        if (
            values
            and basic in {"amount", "number"}
            or (basic == "digits" and AMOUNT_HEADER.search(self.name))
        ):
            detail = []
            if self.amount_styles:
                detail.append("style " + ", ".join(sorted(self.amount_styles)))
            if self.symbols:
                detail.append("symbols " + "".join(sorted(self.symbols)))
            if self.negatives:
                detail.append("negatives as " + ", ".join(sorted(self.negatives)))
            detail.append(f"up to {self.max_decimals} decimal place(s)")
            suggestions.append("possible amount (" + "; ".join(detail) + ")")
            if len(self.amount_styles) > 1:
                notes.append("mixed thousands/decimal separator styles")
            if self.max_decimals > 2:
                notes.append("values with more than 2 decimal places")
        if self.kinds["email"]:
            suggestions.append("possible personal data (email addresses)")
        if self.kinds["phone"]:
            suggestions.append("possible personal data (phone numbers)")

        if self.leading_zeros:
            notes.append(
                f"{self.leading_zeros} value(s) with leading zeros; keep them by "
                "treating this column as text identifiers"
            )
        if self.kinds["number cell"] and id_shaped == values and suggestions:
            notes.append(
                "stored as numeric cells; leading zeros, if there were any, may "
                "already have been removed by the spreadsheet"
            )
        if self.long_numbers:
            notes.append(
                f"{self.long_numbers} number(s) longer than {EXCEL_DIGITS} digits; "
                "Excel cannot hold these exactly as numbers"
            )
        if self.padded:
            notes.append(f"{self.padded} value(s) with surrounding whitespace")

        return {
            "position": self.position,
            "name": self.name,
            "basic_type": basic,
            "values": values,
            "blank": self.blanks,
            "distinct": distinct,
            "kinds": kinds,
            "suggestions": suggestions,
            "notes": notes,
        }


def shape(value, limit=24):
    """
    redacts a cell to its character shape: letters become A/a, digits 9
    :param value: raw cell value
    :param limit: maximum characters shown
    :returns: redacted text
    """
    if is_blank(value):
        return ""
    if isinstance(value, (date, datetime)):
        return "<date cell>"
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return "<number cell>"
    text = str(value).strip()
    redacted = "".join(
        "9"
        if c.isdigit()
        else "A"
        if c.isupper()
        else "a"
        if c.isalpha()
        else c
        if c in " -_./:#@$€£¥,()+'&"
        else "*"
        for c in text[:limit]
    )
    return redacted + ("…" if len(text) > limit else "")


def clean_header(cell):
    """
    trims a header cell the same way the import adapters do
    :param cell: raw header value
    :returns: header text
    """
    if cell is None:
        return ""
    return str(cell).replace("﻿", "").strip()


def analyze_rows(header_cells, rows, max_rows, preview_rows):
    """
    scans the rows of one table
    :param header_cells: raw header row
    :param rows: iterator over the remaining rows
    :param max_rows: maximum data rows to scan
    :param preview_rows: number of redacted rows to keep
    :returns: table report dictionary
    """
    header = [clean_header(cell) for cell in header_cells]
    while header and not header[-1]:
        header.pop()
    warnings = []

    duplicates = sorted({name for name in header if name and header.count(name) > 1})
    if duplicates:
        warnings.append(
            "duplicate column headers "
            + ", ".join(repr(name) for name in duplicates)
            + "; a profile cannot map these until they are renamed"
        )
    blank_positions = [index for index, name in enumerate(header, 1) if not name]
    if blank_positions:
        warnings.append(f"blank header in column(s) {blank_positions}")
    if len(header) == 1 and any(sep in header[0] for sep in (";", "\t", "|")):
        warnings.append(
            "only one column was found and the header contains ';', tab, or '|'; "
            "the file may use a different delimiter, and only comma-separated "
            "CSV is supported"
        )
    if header and all(
        DIGITS.match(name) or AMOUNT_US.match(name) for name in header if name
    ):
        warnings.append("the first row looks like data; a header row is required")

    columns = [
        ColumnStats(name or f"(column {i})", i) for i, name in enumerate(header, 1)
    ]
    scanned = 0
    ragged = 0
    blank_rows = 0
    truncated = False
    preview = []

    for cells in rows:
        cells = list(cells or ())
        if all(is_blank(cell) for cell in cells):
            blank_rows += 1
            continue
        if scanned >= max_rows:
            truncated = True
            break
        scanned += 1
        extra = cells[len(header) :]
        if len(cells) < len(header) or any(not is_blank(c) for c in extra):
            ragged += 1
        for column, cell in zip(columns, cells):
            column.observe(cell)
        for column in columns[len(cells) :]:
            column.observe(None)
        if len(preview) < preview_rows:
            preview.append([shape(cell) for cell in cells[: len(header)]])

    if not header:
        warnings.append("no header row")
    elif scanned == 0:
        warnings.append("header only; no data rows")
    if ragged:
        warnings.append(
            f"{ragged} row(s) have a different number of cells than the header"
        )
    for column in columns:
        if column.values == 0 and scanned:
            warnings.append(f"column {column.name!r} is entirely blank")

    table = {
        "columns": [column.summary(scanned) for column in columns],
        "data_rows": scanned,
        "rows_truncated": truncated,
        "blank_rows_skipped": blank_rows,
        "warnings": warnings,
    }
    if preview_rows:
        table["preview"] = {"header": header, "rows": preview}
    return table


def inspect_csv(path, encoding, max_rows, preview_rows):
    """
    inspects a CSV export
    :param path: file path
    :param encoding: text encoding to decode with
    :param max_rows: maximum data rows to scan
    :param preview_rows: number of redacted rows to keep
    :returns: tuple of table reports and error messages
    """
    try:
        with path.open(newline="", encoding=encoding) as handle:
            reader = csv.reader(handle)
            header = next(reader, None)
            if header is None or not any(cell.strip() for cell in header):
                return [], ["file is empty or its first row is blank"]
            table = analyze_rows(header, reader, max_rows, preview_rows)
    except UnicodeDecodeError as error:
        return [], [
            f"not valid {encoding} text (byte offset {error.start}); the export "
            "may use another encoding such as cp1252, which can be declared with "
            "input.csv_encoding in a profile"
        ]
    except csv.Error as error:
        return [], [f"could not be parsed as CSV: {error}"]

    table["name"] = None
    table["encoding"] = encoding
    return [table], []


def inspect_xlsx(path, max_rows, preview_rows):
    """
    inspects every worksheet of an XLSX workbook
    :param path: file path
    :param max_rows: maximum data rows to scan per worksheet
    :param preview_rows: number of redacted rows to keep
    :returns: tuple of table reports and error messages
    """
    if not zipfile.is_zipfile(path):
        return [], [
            "not a valid .xlsx workbook (password-protected and legacy .xls files "
            "also look like this); re-save as an unprotected .xlsx"
        ]
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises many unrelated exception types
        return [], [f"could not be opened as XLSX: {type(error).__name__}"]

    tables = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None or all(is_blank(cell) for cell in header):
                tables.append(
                    {
                        "name": sheet.title,
                        "columns": [],
                        "data_rows": 0,
                        "rows_truncated": False,
                        "blank_rows_skipped": 0,
                        "warnings": ["worksheet is empty or its first row is blank"],
                    }
                )
                continue
            table = analyze_rows(header, rows, max_rows, preview_rows)
            tables.append({"name": sheet.title, **table})
    except Exception as error:  # corrupt sheet XML surfaces while iterating
        return [], [f"worksheet could not be read: {type(error).__name__}"]
    finally:
        workbook.close()

    return tables, []


def inspect_file(path, encoding="utf-8-sig", max_rows=DEFAULT_MAX_ROWS, preview_rows=0):
    """
    inspects one export file without modifying it
    :param path: file path
    :param encoding: text encoding for CSV files
    :param max_rows: maximum data rows to scan per table
    :param preview_rows: number of redacted rows to keep per table
    :returns: file report dictionary
    """
    path = Path(path)
    suffix = path.suffix.lower()
    report = {
        "file": path.name,
        "format": SUPPORTED_SUFFIXES.get(suffix),
        "bytes": None,
        "tables": [],
        "errors": [],
    }

    try:
        size = path.stat().st_size
    except OSError as error:
        report["errors"].append(f"cannot be read: {error.strerror}")
        return report
    report["bytes"] = size

    if suffix not in SUPPORTED_SUFFIXES:
        report["errors"].append(
            UNSUPPORTED_SUFFIXES.get(suffix, "unsupported file type; use CSV or XLSX")
        )
        return report
    if size == 0:
        report["errors"].append("file is empty")
        return report
    if size > MAX_FILE_BYTES:
        report["errors"].append(
            f"file is larger than the {MAX_FILE_BYTES // (1024 * 1024)} MB "
            "inspection limit"
        )
        return report

    try:
        if suffix == ".csv":
            tables, errors = inspect_csv(path, encoding, max_rows, preview_rows)
        else:
            tables, errors = inspect_xlsx(path, max_rows, preview_rows)
    except OSError as error:
        tables, errors = [], [f"cannot be read: {error.strerror}"]

    report["tables"] = tables
    report["errors"] = errors
    return report


def inspect_path(path, encoding="utf-8-sig", max_rows=DEFAULT_MAX_ROWS, preview_rows=0):
    """
    inspects a file or every file directly inside a directory
    :param path: file or directory path
    :param encoding: text encoding for CSV files
    :param max_rows: maximum data rows to scan per table
    :param preview_rows: number of redacted rows to keep per table
    :returns: list of file reports in name order
    """
    codecs.lookup(encoding)
    path = Path(path)
    if path.is_dir():
        files = sorted(
            (item for item in path.iterdir() if item.is_file()),
            key=lambda item: item.name,
        )
    elif path.is_file():
        files = [path]
    else:
        raise FileNotFoundError(path)

    return [
        inspect_file(item, encoding, max_rows, min(preview_rows, MAX_PREVIEW_ROWS))
        for item in files
        if not item.name.startswith(".")
    ]


def format_report(reports, max_rows):
    """
    renders inspection reports as plain text
    :param reports: list of file reports
    :param max_rows: scan limit used
    :returns: text
    """
    lines = [
        "Everything below is a suggestion inferred from the file contents. Confirm "
        "each mapping with the business before writing a profile.",
    ]
    if not reports:
        lines.append("")
        lines.append("no files found")

    for report in reports:
        lines.append("")
        kind = (report["format"] or "unsupported").upper()
        size = "" if report["bytes"] is None else f", {report['bytes']:,} bytes"
        lines.append(f"{report['file']} ({kind}{size})")
        for error in report["errors"]:
            lines.append(f"  ERROR: {error}")
        if report["format"] == "xlsx" and report["tables"]:
            names = ", ".join(repr(table["name"]) for table in report["tables"])
            lines.append(f"  worksheets: {names}")

        for table in report["tables"]:
            indent = "  "
            if table["name"] is not None:
                lines.append(f"  worksheet {table['name']!r}")
                indent = "    "
            rows = f"{table['data_rows']:,} data row(s)"
            if table["rows_truncated"]:
                rows = f"first {max_rows:,} data rows scanned (more not read)"
            lines.append(f"{indent}{rows}, {len(table['columns'])} column(s)")
            for warning in table["warnings"]:
                lines.append(f"{indent}warning: {warning}")
            for column in table["columns"]:
                distinct = (
                    ""
                    if column["distinct"] is None
                    else f", {column['distinct']} distinct"
                )
                lines.append(
                    f"{indent}{column['position']:>3}. {column['name']}: "
                    f"{column['basic_type']} ({column['values']} values, "
                    f"{column['blank']} blank{distinct})"
                )
                for suggestion in column["suggestions"]:
                    lines.append(f"{indent}       suggestion: {suggestion}")
                for note in column["notes"]:
                    lines.append(f"{indent}       note: {note}")
            preview = table.get("preview")
            if preview:
                lines.append(f"{indent}redacted preview (letters as A/a, digits as 9):")
                lines.append(f"{indent}  " + " | ".join(preview["header"]))
                for row in preview["rows"]:
                    lines.append(f"{indent}  " + " | ".join(row))

    return "\n".join(lines)
