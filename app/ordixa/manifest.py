import csv
import hashlib
import json
import secrets
import shutil
import subprocess
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.ordixa import __version__
from app.ordixa.adapters import file_sha256
from app.ordixa.issues import OUTCOME_LABELS
from app.ordixa.pipeline import LEDGER_COLUMNS, severity_counts
from app.ordixa.workbook import safe_text, write_workbook

MANIFEST_VERSION = 1
PROJECT_ROOT = Path(__file__).resolve().parents[2]
WORKBOOK_NAME = "ordixa_exceptions.xlsx"
LEDGER_NAME = "reconciliation_ledger.csv"
ISSUES_NAME = "rejected_rows.csv"
MANIFEST_NAME = "manifest.json"
ISSUE_COLUMNS = [
    "dataset",
    "file",
    "row",
    "record_id",
    "outcome",
    "prevented_processing",
    "code",
    "field",
    "value",
    "message",
]


def new_run_id(now):
    """
    creates a sortable, collision-resistant run identifier
    :param now: run start time in UTC
    :returns: identifier such as 20261015T183000Z-3fa2c1
    """
    return f"{now:%Y%m%dT%H%M%SZ}-{secrets.token_hex(3)}"


def code_version():
    """
    describes the application version and git commit when available
    :returns: dictionary with version, commit, and dirty flag
    """
    version = {"ordixa": __version__, "git_commit": None, "git_dirty": None}
    try:
        commit = subprocess.run(
            ["git", "rev-parse", "HEAD"],
            cwd=PROJECT_ROOT,
            capture_output=True,
            text=True,
            timeout=5,
            check=True,
        ).stdout.strip()
        status = subprocess.run(
            ["git", "status", "--porcelain", "--untracked-files=no"],
            cwd=PROJECT_ROOT,
            capture_output=True,
            text=True,
            timeout=5,
            check=True,
        ).stdout.strip()
    except OSError, subprocess.SubprocessError:
        return version

    version["git_commit"] = commit or None
    version["git_dirty"] = bool(status)
    return version


def describe_version(version):
    """
    formats a code version for display
    :param version: dictionary from code_version
    :returns: text such as 0.1.0 (a1b2c3d4e5f6)
    """
    commit = version.get("git_commit")
    if not commit:
        return version["ordixa"]
    dirty = ", uncommitted changes" if version.get("git_dirty") else ""
    return f"{version['ordixa']} ({commit[:12]}{dirty})"


def plain(value):
    """
    converts values into JSON and CSV friendly text without losing precision
    :param value: python value
    :returns: serializable value
    """
    if value is None:
        return None
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, dict):
        return {key: plain(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [plain(item) for item in value]
    if not isinstance(value, (str, bool, int, float)) and pd.isna(value):
        return None
    return value


def csv_cell(value):
    """
    formats a value for a CSV artifact, neutralizing formula-like text
    :param value: python value
    :returns: text cell
    """
    value = plain(value)
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return str(value)
    return safe_text(value)


def write_csv(path, columns, rows):
    """
    writes rows to a UTF-8 CSV with LF line endings
    :param path: destination path
    :param columns: column names
    :param rows: list of dictionaries
    :returns: none
    """
    with Path(path).open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(columns)
        for row in rows:
            writer.writerow([csv_cell(row.get(column)) for column in columns])


def issue_rows(issues):
    """
    converts issues to rows for the CSV artifact; original row values are left
    out so the CSV carries no more customer data than the reason requires
    :param issues: list of issues
    :returns: list of dictionaries
    """
    return [
        {
            "dataset": issue.dataset,
            "file": issue.file,
            "row": issue.row,
            "record_id": issue.record_id,
            "outcome": OUTCOME_LABELS[issue.severity],
            "prevented_processing": issue.blocking,
            "code": issue.code,
            "field": issue.field,
            "value": issue.value,
            "message": issue.message,
        }
        for issue in issues
    ]


def result_fingerprint(result):
    """
    hashes the business outcome of a run independently of when it ran
    :param result: run result
    :returns: SHA-256 hex digest
    """
    payload = {
        "ledger": plain(result.ledger.to_dict(orient="records")),
        "issues": issue_rows(result.issues),
        "summary": plain(result.summary),
    }
    encoded = json.dumps(payload, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def run_status(result):
    """
    names the processing status recorded in the manifest
    :param result: run result
    :returns: completed or completed_with_rejections
    """
    return (
        "completed_with_rejections" if result.summary["rejected_rows"] else "completed"
    )


def create_run_directory(out_dir, run_id, staging=False):
    """
    creates a fresh directory for one run; existing runs are never overwritten
    :param out_dir: parent output directory
    :param run_id: run identifier
    :param staging: create the hidden .run-<id>.partial directory instead
    :returns: path of the new run directory
    """
    name = f".run-{run_id}.partial" if staging else f"run-{run_id}"
    run_dir = Path(out_dir) / name
    run_dir.mkdir(parents=True, exist_ok=False)
    return run_dir


def publish_run(out_dir, started_at, write):
    """
    writes a run into a hidden staging directory and publishes it atomically

    the staging directory is renamed to run-<id> only after write() returns,
    so a run-<id> directory never holds partial outputs; the staging
    directory is removed if writing fails
    :param out_dir: parent output directory
    :param started_at: UTC start time
    :param write: function(staging_dir, final_name, run_id) returning a manifest
    :returns: tuple of run directory and manifest dictionary
    """
    run_id = new_run_id(started_at)
    final_dir = Path(out_dir) / f"run-{run_id}"
    if final_dir.exists():
        raise FileExistsError(f"run directory already exists: {final_dir}")
    staging_dir = create_run_directory(out_dir, run_id, staging=True)

    try:
        manifest = write(staging_dir, final_dir.name, run_id)
        staging_dir.rename(final_dir)
    except BaseException:
        shutil.rmtree(staging_dir, ignore_errors=True)
        raise

    return final_dir, manifest


def write_run_outputs(result, out_dir, started_at=None):
    """
    writes the workbook, CSV artifacts, and manifest for a processed run
    :param result: run result
    :param out_dir: parent output directory
    :param started_at: optional UTC start time
    :returns: tuple of run directory and manifest dictionary
    """
    started_at = started_at or datetime.now(timezone.utc)
    return publish_run(
        out_dir,
        started_at,
        lambda run_dir, final_name, run_id: write_artifacts(
            result, run_dir, final_name, run_id, started_at
        ),
    )


def output_entries(run_dir, files):
    """
    describes written artifacts for the manifest
    :param run_dir: directory holding the artifacts
    :param files: (file name, deterministic) pairs in manifest order
    :returns: list of output dictionaries
    """
    return [
        {
            "file": name,
            "sha256": file_sha256(run_dir / name),
            "bytes": (run_dir / name).stat().st_size,
            "deterministic": deterministic,
        }
        for name, deterministic in files
    ]


def write_manifest(run_dir, manifest):
    """
    writes manifest.json with stable formatting
    :param run_dir: directory receiving the manifest
    :param manifest: manifest dictionary
    :returns: none
    """
    (run_dir / MANIFEST_NAME).write_text(
        json.dumps(manifest, indent=2, sort_keys=False) + "\n",
        encoding="utf-8",
        newline="\n",
    )


def write_artifacts(result, run_dir, final_name, run_id, started_at):
    """
    writes every artifact and the manifest into a directory
    :param result: run result
    :param run_dir: directory receiving the files
    :param final_name: directory name recorded in the manifest
    :param run_id: run identifier
    :param started_at: UTC start time
    :returns: manifest dictionary
    """
    version = code_version()
    status = run_status(result)

    write_workbook(
        result,
        run_dir / WORKBOOK_NAME,
        {
            "run_id": run_id,
            "generated_at": started_at.strftime("%Y-%m-%d %H:%M:%S"),
            "code_version": describe_version(version),
            "status": status.replace("_", " "),
        },
    )
    write_csv(
        run_dir / LEDGER_NAME, LEDGER_COLUMNS, result.ledger.to_dict(orient="records")
    )
    write_csv(run_dir / ISSUES_NAME, ISSUE_COLUMNS, issue_rows(result.issues))

    outputs = output_entries(
        run_dir, ((LEDGER_NAME, True), (ISSUES_NAME, True), (WORKBOOK_NAME, False))
    )
    profile = result.profile
    manifest = {
        "manifest_version": MANIFEST_VERSION,
        "result": {
            "status": status,
            "profile": {
                "name": profile.name,
                "version": profile.version,
                "sha256": profile.sha256,
                "identifiers": dict(sorted(profile.identifiers.items())),
            },
            "code_version": version,
            "inputs": [
                {
                    "role": item.role,
                    "file": item.name,
                    "sha256": item.sha256,
                    "data_rows": item.data_rows,
                }
                for item in result.inputs
            ],
            "ignored_files": len(result.ignored_files),
            "row_counts": result.counts,
            "issue_counts": severity_counts(result.issues),
            "artifact_rows": {
                LEDGER_NAME: len(result.ledger),
                ISSUES_NAME: len(result.issues),
            },
            "totals": plain(
                {key: value for key, value in result.summary.items() if key != "issues"}
            ),
            "currency": profile.amounts.currency,
            "result_fingerprint": result_fingerprint(result),
        },
        "execution": {
            "run_id": run_id,
            "started_at": started_at.isoformat(),
            "finished_at": datetime.now(timezone.utc).isoformat(),
            "reconciled_at": plain(result.reconciled_at),
            "output_directory": final_name,
        },
        "outputs": outputs,
    }

    write_manifest(run_dir, manifest)
    return manifest


def find_inputs(name, locations):
    """
    finds candidate copies of a recorded input among directories and files
    :param name: input file name recorded in the manifest
    :param locations: directories or files supplied for verification
    :returns: list of existing candidate paths
    """
    candidates = []
    for location in map(Path, locations):
        if location.is_dir() and (location / name).is_file():
            candidates.append(location / name)
        elif location.is_file() and location.name == name:
            candidates.append(location)
    return candidates


def verify_run(run_dir, input_dir=None):
    """
    re-checks a finished run directory against its manifest

    recomputes the SHA-256 and size of every listed output, the row counts of
    the CSV artifacts, optionally the SHA-256 of every input file, and confirms
    the workbook holds no formula cells
    :param run_dir: run-<id> directory
    :param input_dir: optional directory holding the original exports, or a
        list of directories and files when the inputs came from several places
    :returns: list of (check description, passed) tuples
    """
    run_dir = Path(run_dir)
    manifest = json.loads((run_dir / MANIFEST_NAME).read_text(encoding="utf-8"))
    checks = []

    for output in manifest["outputs"]:
        path = run_dir / output["file"]
        checks.append(
            (
                f"{output['file']} matches its recorded SHA-256 and size",
                path.is_file()
                and file_sha256(path) == output["sha256"]
                and path.stat().st_size == output["bytes"],
            )
        )

    for name, expected in manifest["result"]["artifact_rows"].items():
        with (run_dir / name).open(newline="", encoding="utf-8") as handle:
            rows = sum(1 for _ in csv.reader(handle)) - 1
        checks.append(
            (f"{name} has the recorded {expected} data rows", rows == expected)
        )

    if input_dir is not None:
        locations = input_dir if isinstance(input_dir, (list, tuple)) else [input_dir]
        for item in manifest["result"]["inputs"]:
            checks.append(
                (
                    f"input {item['file']} is unchanged since the run",
                    any(
                        file_sha256(path) == item["sha256"]
                        for path in find_inputs(item["file"], locations)
                    ),
                )
            )

    # every workbook a run lists is checked; reconciliation runs have one
    workbooks = [o["file"] for o in manifest["outputs"] if o["file"].endswith(".xlsx")]
    for name in workbooks:
        workbook = load_workbook(run_dir / name)
        formulas = sum(
            1
            for sheet in workbook
            for row in sheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
        )
        workbook.close()
        description = (
            "workbook contains no formulas"
            if len(workbooks) == 1
            else f"{name} contains no formulas"
        )
        checks.append((description, formulas == 0))
    return checks
