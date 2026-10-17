"""
the comparison workbook: Summary, Added Records, Removed Records,
Modified Records, and Data Quality Issues

it reuses the reconciliation workbook's palette, cell writer, and print
layout; every value goes through the same formula-injection guard
"""

import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.ordixa.tabular import (
    COMPLETE,
    COMPLETE_WITH_WARNINGS,
    DUPLICATE_REMOVED,
    INCOMPLETE,
    NOTE,
    REJECTED,
)
from app.ordixa.workbook import (
    AMBER,
    BOLD_VALUE_FONT,
    BOTTOM_RULE,
    BRAND_FONT,
    FONT_NAME,
    GRAPHITE,
    HASH_FONT,
    HEADER_FILL,
    HEADER_RULE,
    INTEGER_FORMAT,
    LABEL_FONT,
    SECTION_FONT,
    SEVERITY_STYLES,
    STATUS_STYLES,
    TITLE_FONT,
    VALUE_FONT,
    fit_columns,
    fit_rows,
    safe_text,
    set_print_layout,
    style_header,
    write_cell,
    write_table,
)

SHEET_NAMES = (
    "Summary",
    "Added Records",
    "Removed Records",
    "Modified Records",
    "Data Quality Issues",
)
# tab colors reuse the status chip text colors: green added, red removed,
# teal modified, warm amber for data quality
TAB_COLORS = dict(zip(SHEET_NAMES, (GRAPHITE, "1E4D35", "7A2216", "1B4855", AMBER)))
CHANGE_STYLES = {
    "added": STATUS_STYLES["paid"][1:],
    "removed": STATUS_STYLES["unpaid"][1:],
    "modified": STATUS_STYLES["overpaid"][1:],
    "attention": SEVERITY_STYLES["warning"],
}
ASSESSMENT_STYLES = {
    COMPLETE: ("Complete", *STATUS_STYLES["paid"][1:]),
    COMPLETE_WITH_WARNINGS: ("Complete with warnings", *SEVERITY_STYLES["warning"]),
    INCOMPLETE: ("Incomplete", *SEVERITY_STYLES["error"]),
}
ASSESSMENT_GUIDANCE = {
    COMPLETE: "Every record was classified and every configured column compared.",
    COMPLETE_WITH_WARNINGS: (
        "Every record was classified; review the warnings below before relying "
        "on the categories."
    ),
    INCOMPLETE: (
        "Some records or columns could not be compared; Unchanged does not cover "
        "them. Fix the findings below and rerun."
    ),
}
OUTCOME_STYLES = {
    REJECTED: SEVERITY_STYLES["error"],
    DUPLICATE_REMOVED: SEVERITY_STYLES["excluded"],
    NOTE: SEVERITY_STYLES["excluded"],
}
# data columns are keyed with a prefix so a profile column named, say,
# message never picks up the styling of a report column of that name
DATA_PREFIX = "data:"
ISSUE_COLUMNS = [
    "role",
    "file",
    "row",
    "outcome",
    "code",
    "key",
    "column",
    "value",
    "message",
]
ISSUE_HEADERS = [
    "Export",
    "Source file",
    "Source row",
    "Outcome",
    "Issue code",
    "Business key",
    "Column",
    "Value",
    "Reason",
]


def chip(cell, style):
    """
    tints a cell with a status background and matching bold text
    :param cell: worksheet cell
    :param style: (background, text color) pair
    :returns: none
    """
    fill, color = style
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name=FONT_NAME, color=color, bold=True)


def build_summary_sheet(sheet, result, metadata):
    """
    writes record counts, changes by column, guidance, and processing details
    :param sheet: worksheet
    :param result: compare result
    :param metadata: execution metadata for the run
    :returns: none
    """
    profile = result.profile
    counts = result.counts()
    width = 5
    for brand_row in (1, 2):
        for column in range(1, width + 1):
            sheet.cell(row=brand_row, column=column).fill = HEADER_FILL
    for column in range(1, width + 1):
        sheet.cell(row=2, column=column).border = HEADER_RULE
    sheet["A1"] = "Ordixa"
    sheet["A1"].font = BRAND_FONT
    sheet["A1"].alignment = Alignment(vertical="bottom")
    sheet["A2"] = "Change report"
    sheet["A2"].font = TITLE_FONT
    sheet["A2"].alignment = Alignment(vertical="top")
    sheet.row_dimensions[1].height = 32
    sheet.row_dimensions[2].height = 22
    sheet["A3"] = safe_text(
        f"{result.before.input.name} compared with {result.after.input.name} - "
        f"matched by {', '.join(profile.key)} - generated "
        f"{metadata['generated_at']} UTC"
    )
    sheet["A3"].font = LABEL_FONT
    sheet.row_dimensions[3].height = 20
    row = 5

    def section(title):
        nonlocal row
        sheet.cell(row=row, column=1, value=title).font = SECTION_FONT
        for column in range(1, width + 1):
            sheet.cell(row=row, column=column).border = BOTTOM_RULE
        sheet.row_dimensions[row].height = 20
        row += 1

    def line(label, value, emphasis=False, style=None):
        nonlocal row
        label_cell = sheet.cell(row=row, column=1, value=label)
        label_cell.font = LABEL_FONT
        if style:
            chip(label_cell, style)
        cell = write_cell(sheet, row, 2, value)
        cell.font = BOLD_VALUE_FONT if emphasis else VALUE_FONT
        if isinstance(value, int):
            cell.number_format = INTEGER_FORMAT
            cell.alignment = Alignment(horizontal="right")
        row += 1

    def note(label, text):
        nonlocal row
        sheet.cell(row=row, column=1, value=label).font = LABEL_FONT
        sheet.cell(row=row, column=2, value=text).font = VALUE_FONT
        row += 1

    assessment = metadata["assessment"]
    label, *style = ASSESSMENT_STYLES[assessment["level"]]

    def wide_note(label, text, style=None):
        # findings are sentences; merge B:E and wrap so they stay in the
        # printed area instead of running off the right edge
        nonlocal row
        label_cell = sheet.cell(row=row, column=1, value=label)
        label_cell.font = LABEL_FONT
        if style:
            chip(label_cell, style)
        cell = write_cell(sheet, row, 2, text)
        cell.font = VALUE_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        label_cell.alignment = Alignment(vertical="top")
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=width)
        sheet.row_dimensions[row].height = 15 * max(1, math.ceil(len(text) / 80))
        row += 1

    section("Comparison result")
    line("Result", label, True)
    chip(sheet.cell(row=row - 1, column=2), style)
    wide_note("Meaning", ASSESSMENT_GUIDANCE[assessment["level"]])
    for item in assessment["findings"]:
        if item["level"] == INCOMPLETE:
            wide_note("Incomplete", item["message"], SEVERITY_STYLES["error"])
        else:
            wide_note("Warning", item["message"], SEVERITY_STYLES["warning"])
    row += 1

    section("Record counts")
    line("Original records (rows in earlier export)", counts["before_rows"])
    line("New records (rows in later export)", counts["after_rows"])
    line("Added", counts["added"], True, CHANGE_STYLES["added"])
    line("Removed", counts["removed"], True, CHANGE_STYLES["removed"])
    line("Modified", counts["modified"], True, CHANGE_STYLES["modified"])
    line("Unchanged", counts["unchanged"])
    line(
        "Ambiguous keys (conflicting rows)",
        counts["ambiguous_keys"],
        True,
        CHANGE_STYLES["attention"],
    )
    line(
        "Invalid keys (a row was rejected)",
        counts["invalid_keys"],
        True,
        CHANGE_STYLES["attention"],
    )
    line("Rejected rows without a usable key", counts["unkeyed_rejected_rows"])
    row += 1

    section("Changed values by column")
    for index, header in enumerate(["Column", "Changed values"], start=1):
        style_header(sheet.cell(row=row, column=index, value=header), index == 2)
    row += 1
    for name in result.compared_columns:
        sheet.cell(row=row, column=1, value=name).font = VALUE_FONT
        cell = write_cell(
            sheet, row, 2, counts["changed_values_by_column"].get(name, 0)
        )
        cell.font = VALUE_FONT
        cell.number_format = INTEGER_FORMAT
        cell.alignment = Alignment(horizontal="right", indent=1)
        row += 1
    row += 1

    section("Reading this workbook")
    note("Added / Removed Records", "Business keys found in only one export.")
    note("Modified Records", "One line per changed column, with both values.")
    note("Data Quality Issues", "Rows held back from the comparison, and why.")
    note("Ambiguous key", "Several different rows share it; none was chosen.")
    note("Invalid key", "A row with this key was rejected; fix it and rerun.")
    note("Matching", "Records match by business key only, never by row order.")
    note("Calculations", "Exact decimal arithmetic by Ordixa; no formulas.")
    row += 1

    section("Processing")
    line("Run ID", metadata["run_id"])
    line("Generated (UTC)", metadata["generated_at"])
    line("Profile", f"{profile.name} {profile.version}")
    line("Application version", metadata["code_version"])
    line("Status", metadata["status"])
    line("Compared columns", ", ".join(result.compared_columns) or "none")
    if profile.ignore:
        line("Not compared (profile)", ", ".join(profile.ignore))
    if result.skipped_columns:
        line("Not compared (missing)", ", ".join(result.skipped_columns))
    row += 1

    section("Input files")
    for index, header in enumerate(["Export", "File", "Rows", "SHA-256"], start=1):
        style_header(sheet.cell(row=row, column=index, value=header), index == 3)
    row += 1
    for item in result.inputs:
        values = [item.role, item.name, item.data_rows, item.sha256]
        for index, value in enumerate(values, start=1):
            cell = write_cell(sheet, row, index, value)
            cell.font = HASH_FONT if index == 4 else VALUE_FONT
            if index == 3:
                cell.alignment = Alignment(horizontal="right", indent=1)
        row += 1

    for letter, column_width in zip("ABCDE", (40, 26, 10, 20, 20)):
        sheet.column_dimensions[letter].width = column_width
    sheet.sheet_view.showGridLines = False
    set_print_layout(sheet, landscape=False, tab_color=TAB_COLORS[sheet.title])


def numeric_names(profile):
    """
    lists prefixed data columns whose headers sit above figures
    :param profile: table profile
    :returns: set of prefixed column names
    """
    return {
        DATA_PREFIX + column.name
        for column in profile.columns
        if column.type in ("decimal", "integer", "date")
    }


def build_records_sheet(sheet, result, records):
    """
    writes added or removed records with every cleaned column
    :param sheet: worksheet
    :param result: compare result
    :param records: records to list
    :returns: none
    """
    profile = result.profile
    columns = [DATA_PREFIX + column.name for column in profile.columns]
    rows = [
        {DATA_PREFIX + name: value for name, value in record.values.items()}
        for record in records
    ]
    headers = [column.name for column in profile.columns]
    write_table(sheet, 1, columns, headers, rows, numeric_names(profile))
    fit_columns(sheet, columns)
    set_print_layout(sheet, title_rows="1:1", tab_color=TAB_COLORS[sheet.title])


def build_modified_sheet(sheet, result):
    """
    writes one line per changed column with the previous and new values
    :param sheet: worksheet
    :param result: compare result
    :returns: none
    """
    key_columns = [DATA_PREFIX + name for name in result.profile.key]
    columns = key_columns + ["column", "before", "after", "difference"]
    headers = list(result.profile.key) + [
        "Changed column",
        "Previous value",
        "New value",
        "Change",
    ]
    rows = [
        {
            **dict(zip(key_columns, change.key)),
            "column": change.column,
            "before": change.before,
            "after": change.after,
            "difference": change.difference,
        }
        for change in result.changes
    ]
    write_table(sheet, 1, columns, headers, rows, {"difference"})
    fit_columns(sheet, columns)
    set_print_layout(sheet, title_rows="1:1", tab_color=TAB_COLORS[sheet.title])


def issue_rows(issues):
    """
    converts issues into report rows
    :param issues: list of table issues
    :returns: list of dictionaries keyed by ISSUE_COLUMNS
    """
    return [
        {
            "role": issue.side,
            "file": issue.file,
            "row": issue.row or None,
            "outcome": issue.outcome,
            "code": issue.code,
            "key": issue.key,
            "column": issue.column,
            "value": issue.value,
            "message": issue.message,
        }
        for issue in issues
    ]


def build_issues_sheet(sheet, result):
    """
    writes every rejected, removed, or noted row from both exports
    :param sheet: worksheet
    :param result: compare result
    :returns: none
    """
    rows = issue_rows(result.issues)
    last_row = write_table(sheet, 1, ISSUE_COLUMNS, ISSUE_HEADERS, rows, {"row"})
    outcome_column = ISSUE_COLUMNS.index("outcome") + 1
    for row_number, issue in zip(range(2, last_row + 1), result.issues):
        chip(
            sheet.cell(row=row_number, column=outcome_column),
            OUTCOME_STYLES[issue.outcome],
        )
    fit_columns(sheet, ISSUE_COLUMNS, maximum=70)
    fit_rows(sheet, ISSUE_COLUMNS)
    set_print_layout(sheet, title_rows="1:1", tab_color=TAB_COLORS[sheet.title])


def write_change_report(result, path, metadata):
    """
    writes the comparison workbook
    :param result: compare result
    :param path: destination .xlsx path
    :param metadata: execution metadata (run_id, generated_at, code_version,
        status, assessment)
    :returns: destination path
    """
    workbook = Workbook()
    summary = workbook.active
    summary.title = SHEET_NAMES[0]
    build_summary_sheet(summary, result, metadata)
    build_records_sheet(workbook.create_sheet(SHEET_NAMES[1]), result, result.added)
    build_records_sheet(workbook.create_sheet(SHEET_NAMES[2]), result, result.removed)
    build_modified_sheet(workbook.create_sheet(SHEET_NAMES[3]), result)
    build_issues_sheet(workbook.create_sheet(SHEET_NAMES[4]), result)

    workbook.properties.creator = "Ordixa"
    workbook.properties.title = safe_text(f"Ordixa change report {metadata['run_id']}")
    workbook.save(path)
    return path
