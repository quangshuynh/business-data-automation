import csv
import hashlib
import json
from datetime import datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from ordixa_golden import CLEAN_LEDGER, CLEAN_TOTALS
from ordixa_helpers import EXPORTS_DIR, build_profile, copy_exports

from app.ordixa.manifest import (
    ISSUES_NAME,
    LEDGER_NAME,
    MANIFEST_NAME,
    WORKBOOK_NAME,
    create_run_directory,
    write_run_outputs,
)
from app.ordixa.pipeline import process
from app.ordixa.workbook import SHEET_NAMES, safe_text


def run(input_dir, out_dir):
    """
    processes exports and writes run outputs
    :param input_dir: export directory
    :param out_dir: parent output directory
    :returns: tuple of run directory and manifest
    """
    return write_run_outputs(process(build_profile(), input_dir), out_dir)


def as_decimal(value):
    """
    converts a numeric workbook cell back into an exact cent amount
    :param value: number read from the workbook
    :returns: decimal amount
    """
    return Decimal(repr(value)).quantize(Decimal("0.01"))


def sheet_rows(sheet):
    """
    reads a table sheet into dictionaries keyed by header
    :param sheet: worksheet
    :returns: list of dictionaries
    """
    rows = list(sheet.iter_rows(values_only=True))
    return [dict(zip(rows[0], row)) for row in rows[1:]]


def summary_values(sheet):
    """
    reads label/value pairs from the summary sheet
    :param sheet: summary worksheet
    :returns: dictionary of label to value
    """
    return {
        row[0]: row[1]
        for row in sheet.iter_rows(min_col=1, max_col=2, values_only=True)
        if row[0] is not None
    }


@pytest.fixture(scope="module")
def clean_run(tmp_path_factory):
    """
    writes outputs for the clean fixtures once
    :param tmp_path_factory: pytest temporary directory factory
    :returns: tuple of run directory and manifest
    """
    return run(EXPORTS_DIR / "clean", tmp_path_factory.mktemp("out"))


@pytest.fixture(scope="module")
def problematic_run(tmp_path_factory):
    """
    writes outputs for the problematic fixtures once
    :param tmp_path_factory: pytest temporary directory factory
    :returns: tuple of run directory and manifest
    """
    return run(EXPORTS_DIR / "problematic", tmp_path_factory.mktemp("out"))


def test_run_directory_contains_expected_artifacts(clean_run):
    """
    tests the run directory layout
    :param clean_run: clean run directory and manifest
    :returns: none
    """
    run_dir, manifest = clean_run

    assert run_dir.name == f"run-{manifest['execution']['run_id']}"
    assert sorted(path.name for path in run_dir.iterdir()) == sorted(
        [WORKBOOK_NAME, LEDGER_NAME, ISSUES_NAME, MANIFEST_NAME]
    )


def test_workbook_structure_and_table_features(clean_run):
    """
    tests sheet order, frozen headers, filters, and absence of formulas
    :param clean_run: clean run directory and manifest
    :returns: none
    """
    workbook = load_workbook(clean_run[0] / WORKBOOK_NAME)

    assert workbook.sheetnames == list(SHEET_NAMES)
    for name in SHEET_NAMES[1:]:
        sheet = workbook[name]
        assert sheet.freeze_panes == "B2"
        assert sheet.auto_filter.ref.startswith("A1:")
        assert all(cell.font.bold for cell in sheet[1] if cell.value)
    for sheet in workbook:
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f", f"{sheet.title}!{cell.coordinate}"


def test_summary_sheet_matches_hand_computed_totals(clean_run):
    """
    tests summary totals against the independent golden values
    :param clean_run: clean run directory and manifest
    :returns: none
    """
    values = summary_values(load_workbook(clean_run[0] / WORKBOOK_NAME)["Summary"])

    assert values["Orders reconciled"] == 9
    assert as_decimal(values["Order value"]) == CLEAN_TOTALS["order_value"]
    assert as_decimal(values["Payments received"]) == CLEAN_TOTALS["gross_payments"]
    assert as_decimal(values["Refunds (deducted)"]) == CLEAN_TOTALS["refunds"]
    assert as_decimal(values["Adjustments (signed)"]) == CLEAN_TOTALS["adjustments"]
    assert as_decimal(values["Net received"]) == CLEAN_TOTALS["net_received"]
    assert as_decimal(values["Reconciled to orders"]) == CLEAN_TOTALS["applied"]
    assert as_decimal(values["Outstanding balances"]) == CLEAN_TOTALS["outstanding"]
    assert as_decimal(values["Overpayments"]) == CLEAN_TOTALS["overpaid"]
    assert values["Unpaid orders"] == 3
    assert values["Partially paid orders"] == 2
    assert values["Overpaid orders"] == 1
    assert values["Status"] == "completed"
    assert values["Profile"] == "harbor-lane 1.0.0"


def test_ledger_sheet_uses_numeric_and_date_cells(clean_run):
    """
    tests that money and dates are real spreadsheet values with exact cents
    :param clean_run: clean run directory and manifest
    :returns: none
    """
    sheet = load_workbook(clean_run[0] / WORKBOOK_NAME)["Full Ledger"]
    rows = sheet_rows(sheet)
    labels = {"paid": "Paid", "partial": "Partially paid", "unpaid": "Unpaid"}
    labels["overpaid"] = "Overpaid"

    assert len(rows) == len(CLEAN_LEDGER)
    for row in rows:
        total, paid, status, outstanding, overpaid = CLEAN_LEDGER[row["Order ID"]]
        assert isinstance(row["Order total"], (int, float))
        assert as_decimal(row["Order total"]) == total
        assert as_decimal(row["Net paid"]) == paid
        assert as_decimal(row["Outstanding"]) == outstanding
        assert as_decimal(row["Overpaid"]) == overpaid
        assert row["Status"] == labels[status]
        assert isinstance(row["Order date"], datetime)
    assert sheet["H2"].number_format == "#,##0.00;-#,##0.00"
    assert sheet["C2"].number_format == "yyyy-mm-dd"


def test_exceptions_sheet_lists_orders_needing_review(clean_run):
    """
    tests unpaid, partial, and overpaid orders on the exceptions sheet
    :param clean_run: clean run directory and manifest
    :returns: none
    """
    rows = sheet_rows(load_workbook(clean_run[0] / WORKBOOK_NAME)["Exceptions"])

    assert [(row["Exception"], row["Order"]) for row in rows] == [
        ("Unpaid", "#1003"),
        ("Unpaid", "#1005"),
        ("Unpaid", "#1009"),
        ("Partially paid", "#1002"),
        ("Partially paid", "#1007"),
        ("Overpaid", "#1004"),
    ]
    outstanding = sum(as_decimal(row["Outstanding"]) for row in rows)
    assert outstanding == CLEAN_TOTALS["outstanding"]


def test_rejected_rows_sheet_explains_every_issue(problematic_run):
    """
    tests source locations, outcomes, and blocking flags for issues
    :param problematic_run: problematic run directory and manifest
    :returns: none
    """
    rows = sheet_rows(
        load_workbook(problematic_run[0] / WORKBOOK_NAME)["Rejected Rows"]
    )
    by_code = {}
    for row in rows:
        by_code.setdefault(row["Issue code"], []).append(row)

    assert len(rows) == 29
    assert {row["Prevented processing"] for row in by_code["invalid_email"]} == {"No"}
    refund = by_code["refund_sign"][0]
    assert refund["Source file"] == "payments_2026-09.csv"
    assert refund["Source row"] == 2
    assert refund["Outcome"] == "Rejected"
    assert refund["Prevented processing"] == "Yes"
    assert "Amount=$5.00" in refund["Original values"]
    assert by_code["excluded_status"][0]["Outcome"] == "Excluded by profile"
    assert [row["Reason"] for row in by_code["duplicate_row"]] == [
        "exact duplicate of customers.csv row 7; counted once",
        "exact duplicate of orders_2026-08.csv row 9; counted once",
        "exact duplicate of payments_2026-08.csv row 6; counted once",
    ]


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ('=HYPERLINK("http://example.com")', '\'=HYPERLINK("http://example.com")'),
        ("+1+1", "'+1+1"),
        ("@SUM(A1)", "'@SUM(A1)"),
        ("-2+3", "'-2+3"),
        ("\tcmd", "'\tcmd"),
        ("-$89.00", "-$89.00"),
        ("-12.50", "-12.50"),
        ("Avery Quinn", "Avery Quinn"),
        ("bell\x07char", "bellchar"),
    ],
)
def test_safe_text_neutralizes_formulas(text, expected):
    """
    tests formula injection neutralization and control character removal
    :param text: untrusted text
    :param expected: sanitized text
    :returns: none
    """
    assert safe_text(text) == expected


def test_formula_like_source_values_are_stored_as_text(tmp_path):
    """
    tests that hostile export values cannot become formulas in any artifact
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    input_dir = copy_exports("clean", tmp_path / "in")
    customers = input_dir / "customers.csv"
    customers.write_text(
        customers.read_text("utf-8").replace(
            "Avery,Quinn", '"=HYPERLINK(""http://example.com"",""x"")",Quinn'
        ),
        encoding="utf-8",
    )
    payments = input_dir / "payments_2026-08.csv"
    payments.write_text(
        payments.read_text("utf-8").replace(
            "Payment,Paid,$129.99", "Payment,@SUM(1),$129.99"
        ),
        encoding="utf-8",
    )

    run_dir, _ = run(input_dir, tmp_path / "out")
    workbook = load_workbook(run_dir / WORKBOOK_NAME)

    values = [
        cell
        for sheet in workbook
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
        and ("HYPERLINK" in cell.value or "SUM(" in cell.value)
    ]
    assert values
    for cell in values:
        assert cell.data_type == "s"
        assert not cell.value.startswith(("=", "@", "+"))
    for name in (LEDGER_NAME, ISSUES_NAME):
        with (run_dir / name).open(encoding="utf-8") as handle:
            for row in csv.reader(handle):
                for value in row:
                    assert not value.startswith(("=", "@", "+")), (name, value)


def test_manifest_records_fingerprints_and_counts(clean_run):
    """
    tests manifest structure and checksums
    :param clean_run: clean run directory and manifest
    :returns: none
    """
    run_dir, manifest = clean_run
    on_disk = json.loads((run_dir / MANIFEST_NAME).read_text("utf-8"))
    result = on_disk["result"]

    assert on_disk == json.loads(json.dumps(manifest))
    assert result["status"] == "completed"
    assert result["profile"]["name"] == "harbor-lane"
    assert result["profile"]["version"] == "1.0.0"
    assert result["code_version"]["ordixa"] == "0.1.0"
    assert result["row_counts"]["payments"]["accepted"] == 11
    assert result["issue_counts"] == {"error": 0, "warning": 0, "excluded": 0}
    assert result["totals"]["net_received"] == "3752.23"
    for item in result["inputs"]:
        source = EXPORTS_DIR / "clean" / item["file"]
        assert item["sha256"] == hashlib.sha256(source.read_bytes()).hexdigest()
    for output in on_disk["outputs"]:
        digest = hashlib.sha256((run_dir / output["file"]).read_bytes()).hexdigest()
        assert output["sha256"] == digest
    assert {"run_id", "started_at", "finished_at"} <= set(on_disk["execution"])


def test_manifest_contains_no_customer_details(problematic_run):
    """
    tests that names, emails, and raw row values stay out of the manifest
    :param problematic_run: problematic run directory and manifest
    :returns: none
    """
    text = (problematic_run[0] / MANIFEST_NAME).read_text("utf-8")

    for fragment in ("@example", "Quinn", "Moreno", "555-010", "Original", "#2001"):
        assert fragment not in text
    assert json.loads(text)["result"]["status"] == "completed_with_rejections"


def test_reruns_produce_equivalent_results_in_new_directories(tmp_path, clean_run):
    """
    tests repeatability: same inputs give the same business result and
    deterministic artifacts, while execution metadata differs
    :param tmp_path: pytest temporary directory
    :param clean_run: clean run directory and manifest
    :returns: none
    """
    first_dir, first = run(EXPORTS_DIR / "clean", tmp_path)
    second_dir, second = run(EXPORTS_DIR / "clean", tmp_path)

    assert first_dir != second_dir
    assert first["execution"]["run_id"] != second["execution"]["run_id"]
    assert first["result"] == second["result"]
    assert first["result"] == clean_run[1]["result"]
    first_outputs = {item["file"]: item for item in first["outputs"]}
    for item in second["outputs"]:
        if item["deterministic"]:
            assert item["sha256"] == first_outputs[item["file"]]["sha256"]
    assert first_outputs[WORKBOOK_NAME]["deterministic"] is False


def test_changed_inputs_change_the_result_fingerprint(tmp_path, clean_run):
    """
    tests that the fingerprint reflects business content
    :param tmp_path: pytest temporary directory
    :param clean_run: clean run directory and manifest
    :returns: none
    """
    input_dir = copy_exports("clean", tmp_path / "in")
    orders = input_dir / "orders_2026-09.csv"
    orders.write_text(orders.read_text("utf-8").replace("45.00", "46.00"), "utf-8")

    _, manifest = run(input_dir, tmp_path / "out")

    assert (
        manifest["result"]["result_fingerprint"]
        != clean_run[1]["result"]["result_fingerprint"]
    )


def test_existing_run_directory_is_never_overwritten(tmp_path):
    """
    tests that a run directory collision fails instead of overwriting
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    create_run_directory(tmp_path, "fixed")

    with pytest.raises(FileExistsError):
        create_run_directory(tmp_path, "fixed")
