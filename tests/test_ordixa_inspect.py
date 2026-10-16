"""
tests for the read-only inspect command
"""

import json
import zipfile
from datetime import date

import pytest
from openpyxl import Workbook
from ordixa_helpers import EXPORTS_DIR, write_rows

from app.ordixa.cli import EXIT_INPUT, EXIT_OK, EXIT_USAGE, main
from app.ordixa.inspection import inspect_file, inspect_path

TEXT_IDS = EXPORTS_DIR / "text-ids"


def run_inspect(capsys, *args):
    """
    runs the inspect command and captures its output
    :param capsys: pytest capture fixture
    :param args: extra command line arguments
    :returns: tuple of exit code, stdout, stderr
    """
    code = main(["inspect", *map(str, args)])
    captured = capsys.readouterr()
    return code, captured.out, captured.err


def columns_by_name(report, table=0):
    """
    indexes the columns of one table in a file report
    :param report: file report
    :param table: table index
    :returns: dictionary keyed by column name
    """
    return {column["name"]: column for column in report["tables"][table]["columns"]}


def snapshot(directory):
    """
    records every file with its size and modification time
    :param directory: directory to snapshot
    :returns: dictionary keyed by relative path
    """
    return {
        str(path.relative_to(directory)): (path.stat().st_size, path.stat().st_mtime_ns)
        for path in sorted(directory.rglob("*"))
    }


def test_inspect_suggests_identifier_date_and_amount_columns():
    """
    suggestions name likely identifiers, dates with matching formats, and amounts
    :returns: none
    """
    orders, payments = (
        inspect_file(TEXT_IDS / name)
        for name in ("orders_2026-10.csv", "payments_2026-10.csv")
    )
    columns = columns_by_name(orders)
    payment_columns = columns_by_name(payments)

    assert orders["tables"][0]["data_rows"] == 5
    assert columns["Order Number"]["suggestions"] == [
        "possible identifier (unique in scanned rows)"
    ]
    assert "%Y-%m-%d" in columns["Order Date"]["suggestions"][0]
    assert columns["Grand Total"]["basic_type"] == "amount"
    assert "style 1,234.56" in columns["Grand Total"]["suggestions"][0]
    assert "symbols $" in payment_columns["Gross"]["suggestions"][0]
    assert columns["Customer Number"]["blank"] == 1
    assert any("leading zeros" in note for note in columns["Customer Number"]["notes"])


def test_inspect_flags_personal_data_columns():
    """
    email and phone columns are called out for data minimization
    :returns: none
    """
    columns = columns_by_name(inspect_file(TEXT_IDS / "customers.csv"))

    assert "possible personal data (email addresses)" in columns["Email"]["suggestions"]
    assert "possible personal data (phone numbers)" in columns["Phone"]["suggestions"]


def test_inspect_output_contains_no_cell_values(capsys):
    """
    ordinary output describes columns without echoing customer data
    :param capsys: pytest capture fixture
    :returns: none
    """
    code, out, _ = run_inspect(capsys, "--input", TEXT_IDS)

    assert code == EXIT_OK
    assert "suggestion" in out
    for value in (
        "morgan.vale@example.com",
        "Morgan Vale",
        "555-010-0201",
        "00012345",
        "ORD-2026-00142",
        "PAY_8F3K2",
        "1,200.00",
    ):
        assert value not in out


def test_inspect_is_deterministic_and_read_only(capsys, tmp_path, monkeypatch):
    """
    repeated runs print identical output and create or modify no files
    :param capsys: pytest capture fixture
    :param tmp_path: pytest temporary directory
    :param monkeypatch: pytest monkeypatch fixture
    :returns: none
    """
    monkeypatch.chdir(tmp_path)
    before = snapshot(EXPORTS_DIR)

    first = run_inspect(capsys, "--input", EXPORTS_DIR / "problematic")
    second = run_inspect(capsys, "--input", EXPORTS_DIR / "problematic")

    assert first == second
    assert snapshot(EXPORTS_DIR) == before
    assert list(tmp_path.iterdir()) == []


def test_inspect_json_output(capsys):
    """
    --json prints a machine-readable report with the same structure
    :param capsys: pytest capture fixture
    :returns: none
    """
    code, out, _ = run_inspect(capsys, "--input", TEXT_IDS, "--json")
    document = json.loads(out)

    assert code == EXIT_OK
    assert [item["file"] for item in document["files"]] == [
        "customers.csv",
        "orders_2026-10.csv",
        "payments_2026-10.csv",
    ]


def test_inspect_preview_is_opt_in_and_redacted(capsys):
    """
    the preview shows character shapes only
    :param capsys: pytest capture fixture
    :returns: none
    """
    _, without, _ = run_inspect(capsys, "--input", TEXT_IDS / "customers.csv")
    code, out, _ = run_inspect(
        capsys, "--input", TEXT_IDS / "customers.csv", "--preview", "2"
    )

    assert "preview" not in without
    assert code == EXIT_OK
    assert "99999999 | Aaaaaa Aaaa | aaaaaa.aaaa@aaaaaaa.aaa | 999-999-9999" in out
    assert "Morgan" not in out


def test_inspect_reports_day_month_ambiguity(tmp_path):
    """
    dates that fit both month-first and day-first formats are flagged
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    path = write_rows(
        tmp_path / "payments.csv",
        ["Date", "Amount"],
        [["03/04/2026", "$1.00"], ["05/06/2026", "$2.00"]],
    )

    column = columns_by_name(inspect_file(path))["Date"]

    assert any("ambiguous" in note for note in column["notes"])


def test_inspect_csv_structure_problems(tmp_path):
    """
    duplicate and blank headers, ragged rows, and blank columns are reported
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    path = tmp_path / "orders.csv"
    path.write_text(
        "Order,Total,Order,,Notes\n#1,1.00,x,,\n#2,2.00,y,,\n#3,3.00\n",
        encoding="utf-8",
    )

    warnings = " | ".join(inspect_file(path)["tables"][0]["warnings"])

    assert "duplicate column headers 'Order'" in warnings
    assert "blank header in column(s) [4]" in warnings
    assert "1 row(s) have a different number of cells" in warnings
    assert "'Notes' is entirely blank" in warnings


@pytest.mark.parametrize(
    ("content", "message"),
    [
        (b"", "file is empty"),
        (b"\n\n", "file is empty or its first row is blank"),
    ],
)
def test_inspect_empty_csv_files(tmp_path, capsys, content, message):
    """
    empty files are errors and exit with the input code
    :param tmp_path: pytest temporary directory
    :param capsys: pytest capture fixture
    :param content: file bytes
    :param message: expected diagnostic
    :returns: none
    """
    path = tmp_path / "orders.csv"
    path.write_bytes(content)

    code, out, _ = run_inspect(capsys, "--input", path)

    assert code == EXIT_INPUT
    assert message in out


def test_inspect_header_only_and_semicolon_files(tmp_path):
    """
    a header without rows and a semicolon-delimited file are explained
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    header_only = tmp_path / "a.csv"
    header_only.write_text("Order,Total\n", encoding="utf-8")
    semicolons = tmp_path / "b.csv"
    semicolons.write_text("Order;Total\n#1;1,00\n", encoding="utf-8")

    reports = inspect_path(tmp_path)

    assert "header only; no data rows" in reports[0]["tables"][0]["warnings"]
    assert "different delimiter" in reports[1]["tables"][0]["warnings"][0]


def test_inspect_wrong_encoding_is_diagnosed(tmp_path, capsys):
    """
    a cp1252 export is reported under the default encoding and readable with
    an explicit one
    :param tmp_path: pytest temporary directory
    :param capsys: pytest capture fixture
    :returns: none
    """
    path = tmp_path / "customers.csv"
    path.write_bytes("Customer ID,Name\n1,Café Owner\n".encode("cp1252"))

    code, out, _ = run_inspect(capsys, "--input", path)
    retry, _, _ = run_inspect(capsys, "--input", path, "--encoding", "cp1252")
    bad, _, err = run_inspect(capsys, "--input", path, "--encoding", "nope")

    assert code == EXIT_INPUT
    assert "not valid utf-8-sig text" in out
    assert retry == EXIT_OK
    assert bad == EXIT_USAGE
    assert "unknown encoding" in err


def test_inspect_invalid_and_unsupported_workbooks(tmp_path, capsys):
    """
    corrupt, truncated, and legacy workbooks produce diagnostics, not crashes
    :param tmp_path: pytest temporary directory
    :param capsys: pytest capture fixture
    :returns: none
    """
    (tmp_path / "garbage.xlsx").write_bytes(b"this is not a workbook")
    with zipfile.ZipFile(tmp_path / "truncated.xlsx", "w") as archive:
        archive.writestr("hello.txt", "no workbook parts here")
    (tmp_path / "legacy.xls").write_bytes(b"\xd0\xcf\x11\xe0")

    code, out, _ = run_inspect(capsys, "--input", tmp_path)

    assert code == EXIT_INPUT
    assert "garbage.xlsx (XLSX, 22 bytes)\n  ERROR: not a valid .xlsx workbook" in out
    assert "truncated.xlsx" in out and "could not be opened as XLSX" in out
    assert "legacy Excel 97-2003 workbooks are not supported" in out


def test_inspect_xlsx_lists_worksheets_and_native_types(tmp_path):
    """
    every worksheet is described; native dates, numeric ids, and long numbers
    are noted
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Orders"
    sheet.append(["Order ID", "Placed", "Total", "Card"])
    sheet.append([1001, date(2026, 10, 1), 19.99, 1234567890123456789])
    sheet.append([1002, date(2026, 10, 2), 5, 1234567890123456780])
    workbook.create_sheet("Notes")
    workbook.save(tmp_path / "orders.xlsx")

    report = inspect_file(tmp_path / "orders.xlsx")
    columns = columns_by_name(report)

    assert [table["name"] for table in report["tables"]] == ["Orders", "Notes"]
    assert report["tables"][1]["warnings"] == [
        "worksheet is empty or its first row is blank"
    ]
    assert columns["Placed"]["suggestions"] == ["possible date (2 native date cell(s))"]
    assert any("numeric cells" in note for note in columns["Order ID"]["notes"])
    assert any("longer than 15 digits" in n for n in columns["Card"]["notes"])


def test_inspect_numeric_amount_cells_report_decimals_not_identifier_notes(tmp_path):
    """
    native number cells in a money column count their decimal places and do
    not get the leading-zero note meant for identifier columns
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Transaction ID", "Amount"])
    sheet.append(["PMT-1", 84.5])
    sheet.append(["PMT-2", 27.25])
    sheet.append(["PMT-3", -12])
    workbook.save(tmp_path / "payments.xlsx")

    amount = columns_by_name(inspect_file(tmp_path / "payments.xlsx"))["Amount"]

    assert amount["basic_type"] == "number"
    assert amount["suggestions"] == ["possible amount (up to 2 decimal place(s))"]
    assert not any("numeric cells" in note for note in amount["notes"])


def test_inspect_row_limit_is_reported(tmp_path, capsys):
    """
    scanning stops at --max-rows and says so
    :param tmp_path: pytest temporary directory
    :param capsys: pytest capture fixture
    :returns: none
    """
    path = write_rows(
        tmp_path / "orders.csv",
        ["Order", "Total"],
        [[f"#{i}", "1.00"] for i in range(30)],
    )

    code, out, _ = run_inspect(capsys, "--input", path, "--max-rows", "10")

    assert code == EXIT_OK
    assert "first 10 data rows scanned (more not read)" in out


def test_inspect_missing_path(tmp_path, capsys):
    """
    a missing input path is an input error
    :param tmp_path: pytest temporary directory
    :param capsys: pytest capture fixture
    :returns: none
    """
    code, _, err = run_inspect(capsys, "--input", tmp_path / "nope")

    assert code == EXIT_INPUT
    assert "does not exist" in err
