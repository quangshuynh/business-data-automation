"""
tests for the workbook's presentation: palette contrast, guidance text,
explicit fonts, and print layout
"""

import pytest
from openpyxl import load_workbook
from ordixa_helpers import EXPORTS_DIR, build_profile

from app.ordixa import workbook
from app.ordixa.manifest import WORKBOOK_NAME, write_run_outputs
from app.ordixa.pipeline import process


def relative_luminance(hex_color):
    """
    computes the WCAG relative luminance of an RRGGBB color
    :param hex_color: color such as 111D22
    :returns: luminance between 0 and 1
    """
    channels = []
    for start in (0, 2, 4):
        value = int(hex_color[start : start + 2], 16) / 255
        channels.append(
            value / 12.92 if value <= 0.03928 else ((value + 0.055) / 1.055) ** 2.4
        )
    red, green, blue = channels
    return 0.2126 * red + 0.7152 * green + 0.0722 * blue


def contrast(foreground, background):
    """
    computes the WCAG contrast ratio of two colors
    :param foreground: RRGGBB text color
    :param background: RRGGBB background color
    :returns: ratio from 1 to 21
    """
    light, dark = sorted(
        (relative_luminance(foreground), relative_luminance(background)), reverse=True
    )
    return (light + 0.05) / (dark + 0.05)


@pytest.fixture(scope="module")
def demo_workbook(tmp_path_factory):
    """
    writes the demo run once and opens its workbook
    :param tmp_path_factory: pytest temporary directory factory
    :returns: loaded workbook
    """
    run_dir, _ = write_run_outputs(
        process(build_profile(), EXPORTS_DIR / "demo"), tmp_path_factory.mktemp("out")
    )
    return load_workbook(run_dir / WORKBOOK_NAME)


def test_every_text_color_pair_meets_wcag_aaa_contrast():
    """
    tests that each foreground/background pair reaches a 7:1 ratio
    :returns: none
    """
    pairs = [
        (workbook.OFF_WHITE, workbook.GRAPHITE),
        (workbook.CHARTREUSE, workbook.GRAPHITE),
        (workbook.MUTED, "FFFFFF"),
        (workbook.GRAPHITE, "FFFFFF"),
    ]
    pairs += [(fg, bg) for _, bg, fg in workbook.STATUS_STYLES.values()]
    pairs += [(fg, bg) for bg, fg in workbook.SEVERITY_STYLES.values()]

    for foreground, background in pairs:
        assert contrast(foreground, background) >= 7, (foreground, background)


def test_exceptions_explain_what_to_check(demo_workbook):
    """
    tests that every exception carries plain-language guidance
    :param demo_workbook: demo workbook
    :returns: none
    """
    rows = list(demo_workbook["Exceptions"].iter_rows(values_only=True))
    header = rows[0]
    kind, note = header.index("Exception"), header.index("What to check")

    assert [row[kind] for row in rows[1:]] == [
        "Unpaid",
        "Unpaid",
        "Partially paid",
        "Partially paid",
        "Partially paid",
        "Overpaid",
    ]
    for row in rows[1:]:
        assert row[note] == workbook.REVIEW_NOTES[row[kind]]


def test_summary_explains_outcomes_and_sheets(demo_workbook):
    """
    tests the reading guide on the summary sheet
    :param demo_workbook: demo workbook
    :returns: none
    """
    labels = {
        row[0]: row[1]
        for row in demo_workbook["Summary"].iter_rows(max_col=2, values_only=True)
        if row[0]
    }

    assert demo_workbook["Summary"]["A1"].value == "Ordixa"
    for label in (
        "Exceptions sheet",
        "Rejected Rows sheet",
        "Full Ledger sheet",
        "Rejected",
        "Kept with warning",
        "Excluded by profile",
    ):
        assert labels[label]


def test_styled_text_always_names_its_font(demo_workbook):
    """
    tests that no styled cell falls back to an application default font
    :param demo_workbook: demo workbook
    :returns: none
    """
    for sheet in demo_workbook:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and cell.has_style:
                    assert cell.font.name, f"{sheet.title}!{cell.coordinate}"


def test_sheets_print_one_page_wide_with_repeated_headers(demo_workbook):
    """
    tests the print layout of every sheet
    :param demo_workbook: demo workbook
    :returns: none
    """
    for sheet in demo_workbook:
        assert sheet.page_setup.fitToWidth == 1
        assert sheet.page_setup.fitToHeight == 0
        assert sheet.sheet_properties.pageSetUpPr.fitToPage
        assert sheet.sheet_properties.tabColor is not None
    assert demo_workbook["Summary"].page_setup.orientation == "portrait"
    for name in workbook.SHEET_NAMES[1:]:
        assert demo_workbook[name].page_setup.orientation == "landscape"
        assert demo_workbook[name].print_title_rows == "$1:$1"


def test_wrapped_rows_are_tall_enough(demo_workbook):
    """
    tests that rows holding wrapped guidance text get explicit heights
    :param demo_workbook: demo workbook
    :returns: none
    """
    sheet = demo_workbook["Exceptions"]
    for row in range(2, sheet.max_row + 1):
        assert sheet.row_dimensions[row].height >= 2 * workbook.LINE_HEIGHT
