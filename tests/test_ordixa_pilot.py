"""
acceptance tests for the fictional November 2026 pilot onboarding scenario
"""

import csv
import json
import re
import shutil
import subprocess
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from ordixa_golden import (
    DEMO_LEDGER,
    PILOT_LEDGER,
    PILOT_TOTALS,
    PILOT_UNMATCHED,
    WINDOW_TOTALS,
)
from ordixa_helpers import EXPORTS_DIR, PROFILE_PATH, ROOT, build_profile

from app.ordixa import cli
from app.ordixa.manifest import LEDGER_NAME, MANIFEST_NAME, WORKBOOK_NAME
from app.ordixa.pipeline import process

PILOT_DIR = EXPORTS_DIR / "pilot"
FINAL_DIR = PILOT_DIR / "3-final-delivery"
RESERVED_DOMAINS = {"example.com", "example.net", "example.org"}
EMAIL_LIKE = re.compile(r"[A-Za-z0-9._%+-]+@([A-Za-z0-9.-]+\.[A-Za-z]{2,})")
PHONE_LIKE = re.compile(r"\(?\b\d{3}\)?[ .-]\d{3}[ .-]\d{4}\b")


@pytest.fixture(scope="module")
def pilot_result():
    """
    processes the final pilot delivery once
    :returns: run result
    """
    return process(build_profile(), FINAL_DIR)


def golden_tuple(row):
    """
    converts a ledger row into the golden comparison tuple
    :param row: ledger row dictionary
    :returns: tuple of total, net paid, status, outstanding, overpaid
    """
    return (
        row["order_total"],
        row["net_paid"],
        row["financial_status"],
        row["outstanding_balance"],
        row["overpayment_amount"],
    )


def plain_amount(text):
    """
    converts a fixture amount to a decimal with deliberately minimal parsing
    :param text: amount text such as -$1,075.00
    :returns: decimal amount
    """
    return Decimal(text.replace("$", "").replace(",", ""))


def test_final_delivery_matches_hand_computed_ledger(pilot_result):
    """
    tests every reconciled pilot order against independent expectations
    :param pilot_result: processed pilot batch
    :returns: none
    """
    ledger = {
        row["order_id"]: golden_tuple(row)
        for row in pilot_result.ledger.to_dict(orient="records")
    }
    assert ledger == PILOT_LEDGER
    for key, expected in PILOT_TOTALS.items():
        assert pilot_result.summary[key] == expected, key


def test_unmatched_payments_are_rejected_not_guessed(pilot_result):
    """
    tests that payments for orders outside the batch are listed, not matched
    :param pilot_result: processed pilot batch
    :returns: none
    """
    rejected = {
        issue.row: (issue.record_id, issue.value, issue.code)
        for issue in pilot_result.issues
    }
    assert rejected == {
        row: (record, order, "order_not_in_run")
        for row, (record, order, _) in PILOT_UNMATCHED.items()
    }
    # the #4016 payment equals the #4106 total but must not be matched by amount
    refs = ",".join(pilot_result.ledger["transaction_refs"])
    assert "PMT-291106" not in refs
    assert "PMT-291107" not in refs


def test_raw_pilot_files_agree_with_expectations_without_app_parsers():
    """
    tests the source CSVs against the expectations using minimal parsing
    :returns: none
    """
    with (FINAL_DIR / "orders_2026-11.csv").open(newline="", encoding="utf-8") as f:
        orders = {
            int(row["Name"][1:]): plain_amount(row["Total"])
            for row in csv.DictReader(f)
        }
    assert orders == {key: value[0] for key, value in PILOT_LEDGER.items()}
    assert sum(orders.values()) == PILOT_TOTALS["order_value"]

    totals = {"Payment": Decimal(0), "Refund": Decimal(0)}
    with (FINAL_DIR / "payments_2026-11.csv").open(newline="", encoding="utf-8") as f:
        for row_number, row in enumerate(csv.DictReader(f), 2):
            if row_number in PILOT_UNMATCHED:
                assert plain_amount(row["Amount"]) == PILOT_UNMATCHED[row_number][2]
                continue
            totals[row["Type"]] += plain_amount(row["Amount"])
    assert totals["Payment"] == PILOT_TOTALS["payments_received"]
    assert -totals["Refund"] == PILOT_TOTALS["refunds"]
    assert sum(totals.values()) == PILOT_TOTALS["net_received"]


@pytest.mark.parametrize(
    ("delivery", "message"),
    [
        ("1-first-delivery", "no payments file found"),
        ("2-second-delivery", "'Order Reference' (found 'Order Ref')"),
    ],
)
def test_incomplete_deliveries_fail_safely(delivery, message, tmp_path, capsys):
    """
    tests that an incomplete delivery stops the run and leaves no outputs
    :param delivery: pilot delivery folder
    :param message: expected part of the error
    :param tmp_path: pytest temporary directory
    :param capsys: pytest output capture
    :returns: none
    """
    out = tmp_path / "runs"
    exit_code = cli.main(
        [
            "run",
            "--profile",
            str(PROFILE_PATH),
            "--input",
            str(PILOT_DIR / delivery),
            "--out",
            str(out),
        ]
    )

    assert exit_code == cli.EXIT_INPUT
    assert message in capsys.readouterr().err
    assert not out.exists() or not any(out.iterdir())


def test_final_delivery_run_writes_verified_outputs(tmp_path, capsys):
    """
    tests the full operator path: run, manifest totals, workbook, verify
    :param tmp_path: pytest temporary directory
    :param capsys: pytest output capture
    :returns: none
    """
    out = tmp_path / "runs"
    args = ["--profile", str(PROFILE_PATH), "--input", str(FINAL_DIR)]
    assert cli.main(["run", *args, "--out", str(out)]) == cli.EXIT_OK
    [run_dir] = list(out.iterdir())

    manifest = json.loads((run_dir / MANIFEST_NAME).read_text(encoding="utf-8"))
    assert manifest["result"]["status"] == "completed_with_rejections"
    totals = manifest["result"]["totals"]
    assert Decimal(totals["net_received"]) == PILOT_TOTALS["net_received"]
    assert Decimal(totals["outstanding_balance"]) == PILOT_TOTALS["outstanding_balance"]

    with (run_dir / LEDGER_NAME).open(newline="", encoding="utf-8") as handle:
        ledger = {int(row["order_id"]): row for row in csv.DictReader(handle)}
    assert {key: row["financial_status"] for key, row in ledger.items()} == {
        key: value[2] for key, value in PILOT_LEDGER.items()
    }

    workbook = load_workbook(run_dir / WORKBOOK_NAME, read_only=True)
    exceptions = {
        cell
        for row in workbook["Exceptions"].iter_rows(values_only=True)
        for cell in row
        if isinstance(cell, str) and re.fullmatch(r"#\d{4}", cell)
    }
    workbook.close()
    assert exceptions == {"#4102", "#4103", "#4104", "#4106"}

    capsys.readouterr()
    verify = ["verify", "--run", str(run_dir), "--input", str(FINAL_DIR)]
    assert cli.main(verify) == cli.EXIT_OK


def test_bounded_window_matches_an_earlier_period_order(tmp_path):
    """
    tests that supplying both periods' orders and payments in one run matches
    a late payment, as described in the cross-run assessment
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    for source in (
        EXPORTS_DIR / "demo" / "orders_2026-10.csv",
        EXPORTS_DIR / "demo" / "payments_2026-10.xlsx",
        FINAL_DIR / "orders_2026-11.csv",
        FINAL_DIR / "payments_2026-11.csv",
    ):
        shutil.copy(source, tmp_path / source.name)

    result = process(build_profile(), tmp_path)
    ledger = {
        row["order_id"]: golden_tuple(row)
        for row in result.ledger.to_dict(orient="records")
    }

    total = DEMO_LEDGER[3004][0]
    assert ledger[3004] == (total, total, "paid", Decimal("0.00"), Decimal("0.00"))
    assert ledger[4106] == PILOT_LEDGER[4106]
    for key, expected in WINDOW_TOTALS.items():
        assert result.summary[key] == expected, key
    unmatched = sorted(
        issue.record_id for issue in result.issues if issue.code == "order_not_in_run"
    )
    assert unmatched == ["PMT-281015", "PMT-291107"]


def test_every_example_export_is_fictional():
    """
    tests that all example exports, including the pilot set, use reserved
    email domains and the fictional 555-010 phone range
    :returns: none
    """
    paths = sorted(EXPORTS_DIR.rglob("*.csv"))
    assert any("pilot" in path.parts for path in paths)
    for path in paths:
        text = path.read_text(encoding="utf-8")
        for domain in EMAIL_LIKE.findall(text):
            assert domain.lower() in RESERVED_DOMAINS, f"{path.name}: {domain}"
        for phone in PHONE_LIKE.findall(text):
            assert re.sub(r"\D", "", phone)[3:6] == "010", f"{path.name}: {phone}"


def test_tracked_files_contain_no_real_email_addresses():
    """
    tests that no tracked text file carries an email outside reserved domains
    :returns: none
    """
    try:
        listed = subprocess.run(
            ["git", "ls-files", "-z"],
            cwd=ROOT,
            capture_output=True,
            check=True,
        ).stdout
    except OSError, subprocess.CalledProcessError:
        pytest.skip("git is not available")

    for name in filter(None, listed.decode().split("\0")):
        path = ROOT / name
        if path.suffix.lower() in {".png", ".xlsx"} or not path.is_file():
            continue
        text = path.read_text(encoding="utf-8", errors="ignore")
        for domain in EMAIL_LIKE.findall(text):
            # inspect previews redact addresses to shapes such as aaaa@aaaaaaa.aaa
            redacted = set(domain) <= set("a.")
            assert domain.lower() in RESERVED_DOMAINS or redacted, f"{name}: {domain}"
