"""
tests that a comparison never presents missing or unclassified data as
unchanged without saying so

every comparison is assessed as complete, complete with warnings, or
incomplete; these tests pin which situations land in which state
"""

import json
import unicodedata

import pytest
from openpyxl import load_workbook
from table_helpers import (
    COMPARE_PROFILE,
    HEADER,
    INVENTORY_DIR,
    table_profile,
    write_table,
)

from app.ordixa.manifest import MANIFEST_NAME, verify_run
from app.ordixa.tabular import (
    COMPLETE,
    COMPLETE_WITH_WARNINGS,
    INCOMPLETE,
    MODIFIED,
    UNCHANGED,
    compare_tables,
)
from app.ordixa.tabular_outputs import REPORT_NAME, write_compare_run
from app.ordixa.tabular_profile import load_table_profile

ROWS = [
    ["SKU-1", "Mug", "4.50", "10", "2026-01-05"],
    ["SKU-2", "Bowl", "7.00", "3", "2026-01-05"],
    ["SKU-3", "Plate", "9.25", "0", "2026-01-05"],
]
COMPOSED = unicodedata.normalize("NFC", "Café Mug")
DECOMPOSED = unicodedata.normalize("NFD", "Café Mug")


def compare(tmp_path, before_rows, after_rows, profile=None, **headers):
    """
    writes two small exports and compares them
    :param tmp_path: pytest temporary directory
    :param before_rows: rows of the earlier export
    :param after_rows: rows of the later export
    :param profile: optional table profile, the default test profile otherwise
    :param headers: optional before= and after= header overrides
    :returns: compare result
    """
    before = write_table(tmp_path / "before.csv", before_rows, headers.get("before"))
    after = write_table(tmp_path / "after.csv", after_rows, headers.get("after"))
    return compare_tables(profile or table_profile(), before, after)


def codes(assessment):
    """
    lists assessment finding codes
    :param assessment: assessment dictionary
    :returns: list of codes in report order
    """
    return [item["code"] for item in assessment["findings"]]


def test_identical_exports_are_complete(tmp_path):
    """
    tests the clean case: no findings at all
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    result = compare(tmp_path, ROWS, ROWS)
    assert result.assessment() == {"level": COMPLETE, "findings": []}
    assert set(result.status.values()) == {UNCHANGED}


def test_reordered_rows_and_columns_change_nothing(tmp_path):
    """
    tests that matching is by key and header name, never by position
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    order = [4, 2, 0, 3, 1]
    shuffled = [[row[i] for i in order] for row in reversed(ROWS)]
    result = compare(
        tmp_path, ROWS, shuffled, after=[HEADER[i] for i in order], before=HEADER
    )
    assert result.assessment()["level"] == COMPLETE
    assert set(result.status.values()) == {UNCHANGED}


def test_optional_column_missing_from_one_export_is_incomplete(tmp_path):
    """
    tests that a column absent from one export is never silently unchanged
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    after_header = ["SKU", "Name", "Price", "Counted"]
    after = [[row[0], row[1], row[2], row[4]] for row in ROWS]
    after[0][2] = "4.75"
    result = compare(tmp_path, ROWS, after, after=after_header)

    assert result.skipped_columns == ["qty"]
    assert "qty" not in result.compared_columns
    assessment = result.assessment()
    assert assessment["level"] == INCOMPLETE
    (missing,) = [
        f
        for f in assessment["findings"]
        if f["code"] == "column_missing_from_one_export"
    ]
    assert missing["level"] == INCOMPLETE
    assert missing["count"] == 3
    assert "qty" in missing["message"] and "later export" in missing["message"]
    assert "source_headers_differ" in codes(assessment)
    # the price change is still found; only qty is unknown
    assert result.counts()["modified"] == 1


def test_optional_column_missing_from_both_exports_is_a_warning(tmp_path):
    """
    tests that a column neither export has is disclosed but not incomplete
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    header = ["SKU", "Name", "Price", "Counted"]
    rows = [[row[0], row[1], row[2], row[4]] for row in ROWS]
    result = compare(tmp_path, rows, rows, before=header, after=header)
    assessment = result.assessment()
    assert assessment["level"] == COMPLETE_WITH_WARNINGS
    assert codes(assessment) == ["column_missing_from_both"]


def test_ignored_column_differences_are_disclosed(tmp_path):
    """
    tests that ignored columns are not compared, but their differences counted
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    profile = table_profile(mutate=lambda d: d.update(compare={"ignore": ["counted"]}))
    after = [row[:4] + ["2026-02-09"] for row in ROWS]
    result = compare(tmp_path, ROWS, after, profile)

    assert set(result.status.values()) == {UNCHANGED}
    assert result.ignored_differences == {"counted": 3}
    assessment = result.assessment()
    assert assessment["level"] == COMPLETE_WITH_WARNINGS
    (finding,) = assessment["findings"]
    assert finding["code"] == "ignored_column_differs" and finding["count"] == 3


def test_unicode_only_value_change_is_flagged_not_merged(tmp_path):
    """
    tests that NFC and NFD spellings stay a change, with a note explaining it
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    before = [[ROWS[0][0], COMPOSED, *ROWS[0][2:]], *ROWS[1:]]
    after = [[ROWS[0][0], DECOMPOSED, *ROWS[0][2:]], *ROWS[1:]]
    result = compare(tmp_path, before, after)

    assert result.status[("SKU-1",)] == MODIFIED
    notes = [i for i in result.issues if i.code == "unicode_form_differs"]
    assert [(n.key, n.column, n.value) for n in notes] == [("SKU-1", "name", "")]
    assert codes(result.assessment()) == ["unicode_only_changes"]


def test_unicode_only_key_difference_is_flagged_on_both_keys(tmp_path):
    """
    tests that a re-encoded text key is shown as a likely pair, never merged
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    profile = table_profile(
        mutate=lambda d: d["columns"]["sku"].update(type="text"),
    )
    result = compare(
        tmp_path,
        [[COMPOSED, "Mug", "1.00", "1", ""]],
        [[DECOMPOSED, "Mug", "1.00", "1", ""]],
        profile,
    )
    counts = result.counts()
    assert (counts["added"], counts["removed"]) == (1, 1)
    notes = [i for i in result.issues if i.code == "key_unicode_form_differs"]
    assert len(notes) == 2
    assessment = result.assessment()
    assert assessment["level"] == COMPLETE_WITH_WARNINGS
    assert codes(assessment) == ["unicode_key_pairs"]


def test_decimal_spellings_of_the_same_amount_are_unchanged(tmp_path):
    """
    tests exact decimal comparison: $4.50, 4.5, and 4.50 are one amount
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    after = [list(row) for row in ROWS]
    after[0][2] = "$4.5"
    after[1][2] = "7"
    after[2][2] = " 9.25 "
    result = compare(tmp_path, ROWS, after)
    assert set(result.status.values()) == {UNCHANGED}

    after[0][2] = "4.51"
    result = compare(tmp_path, ROWS, after)
    (change,) = result.changes
    assert str(change.difference) == "0.01"


def test_null_empty_and_blank_token_cells_compare_as_empty(tmp_path):
    """
    tests that an empty cell and a blank token are the same blank value
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    before = [row[:4] + [""] for row in ROWS]
    after = [row[:4] + ["N/A"] for row in ROWS]
    result = compare(tmp_path, before, after)
    assert set(result.status.values()) == {UNCHANGED}
    assert result.assessment()["level"] == COMPLETE


def test_blank_token_that_is_a_real_value_is_a_warning(tmp_path):
    """
    tests that a broad blank token such as 0 is reported when it erases values
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    profile = table_profile(blank_values=["N/A", "0"])
    result = compare(tmp_path, ROWS, ROWS, profile)
    assessment = result.assessment()
    assert assessment["level"] == COMPLETE_WITH_WARNINGS
    (finding,) = assessment["findings"]
    assert finding["code"] == "blank_token_is_a_value"
    assert finding["count"] == 2 and "qty" in finding["message"]


def test_rejected_and_conflicting_keys_make_the_comparison_incomplete(tmp_path):
    """
    tests that invalid and ambiguous keys are never counted as unchanged
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    after = [
        ROWS[0],
        ["SKU-2", "Bowl", "not a price", "3", ""],
        ["SKU-2", "Bowl", "7.00", "3", ""],
        ROWS[2],
        ["SKU-3", "Plate", "9.99", "0", ""],
        ["", "Lost", "1.00", "1", ""],
    ]
    result = compare(tmp_path, ROWS, after)
    assert result.status[("SKU-2",)] == "invalid"
    assert result.status[("SKU-3",)] == "ambiguous"
    assessment = result.assessment()
    assert assessment["level"] == INCOMPLETE
    assert codes(assessment)[:2] == ["unclassified_keys", "rejected_rows_without_key"]
    assert all(f["level"] == INCOMPLETE for f in assessment["findings"][:2])


def test_exact_duplicates_are_a_warning_only(tmp_path):
    """
    tests that removing identical rows keeps the comparison complete
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    result = compare(tmp_path, ROWS, [*ROWS, ROWS[1]])
    assessment = result.assessment()
    assert assessment["level"] == COMPLETE_WITH_WARNINGS
    assert codes(assessment) == ["duplicate_rows_removed"]


def test_assessment_messages_never_contain_cell_values(tmp_path):
    """
    tests that findings name columns and counts only
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    secret = "Zebulon Quartzfield"
    before = [[ROWS[0][0], unicodedata.normalize("NFC", secret + "é"), *ROWS[0][2:]]]
    after = [
        [ROWS[0][0], unicodedata.normalize("NFD", secret + "é"), "1.00", "", ""],
        ["SKU-9", secret, "bad", "1", ""],
    ]
    result = compare(tmp_path, before, after)
    text = json.dumps(result.assessment())
    assert "Zebulon" not in text and "SKU-" not in text


@pytest.fixture(scope="module")
def inventory_run(tmp_path_factory):
    """
    writes the fictional inventory comparison once for output tests
    :param tmp_path_factory: pytest temporary directory factory
    :returns: tuple of run directory and manifest
    """
    profile = load_table_profile(COMPARE_PROFILE)
    result = compare_tables(
        profile,
        INVENTORY_DIR / "inventory-before.csv",
        INVENTORY_DIR / "inventory-after.csv",
    )
    return write_compare_run(result, tmp_path_factory.mktemp("compare"))


def test_inventory_comparison_is_assessed_in_the_manifest(inventory_run):
    """
    tests the recorded assessment of the fictional inventory exports
    :param inventory_run: run directory and manifest
    :returns: none
    """
    run_dir, manifest = inventory_run
    assessment = manifest["result"]["assessment"]
    assert assessment["level"] == INCOMPLETE
    assert codes(assessment) == [
        "unclassified_keys",
        "rejected_rows_without_key",
        "ignored_column_differs",
        "duplicate_rows_removed",
    ]
    stored = json.loads((run_dir / MANIFEST_NAME).read_text(encoding="utf-8"))
    assert stored["result"]["assessment"] == assessment
    assert all(passed for _, passed in verify_run(run_dir, [INVENTORY_DIR]))


def test_workbook_summary_leads_with_the_assessment(inventory_run):
    """
    tests that the Summary sheet states the result before the counts
    :param inventory_run: run directory and manifest
    :returns: none
    """
    run_dir, _ = inventory_run
    workbook = load_workbook(run_dir / REPORT_NAME)
    sheet = workbook["Summary"]
    labels = [sheet.cell(row=r, column=1).value for r in range(1, sheet.max_row + 1)]
    values = [sheet.cell(row=r, column=2).value for r in range(1, sheet.max_row + 1)]
    assert labels.index("Comparison result") < labels.index("Record counts")
    assert "Incomplete" in values
    assert labels.count("Incomplete") == 2 and labels.count("Warning") == 2
    # findings wrap inside the printed columns instead of running off the page
    finding_rows = [r for r, label in enumerate(labels, start=1) if label == "Warning"]
    merged = {str(rng) for rng in sheet.merged_cells.ranges}
    for r in finding_rows:
        assert f"B{r}:E{r}" in merged
        assert sheet.cell(row=r, column=2).alignment.wrap_text
    workbook.close()
