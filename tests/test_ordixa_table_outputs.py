"""
tests for clean/compare run directories, manifests, and the change report
"""

import csv
import json
import shutil

import pytest
from openpyxl import load_workbook
from table_helpers import (
    CLEAN_PROFILE,
    COMPARE_PROFILE,
    INVENTORY_DIR,
    table_profile,
    write_table,
)

from app.ordixa import tabular_outputs
from app.ordixa.change_report import SHEET_NAMES, TAB_COLORS
from app.ordixa.manifest import MANIFEST_NAME, verify_run
from app.ordixa.tabular import clean_table, compare_tables
from app.ordixa.tabular_outputs import (
    ADDED_NAME,
    CLEANED_NAME,
    MODIFIED_NAME,
    REPORT_NAME,
    write_clean_run,
    write_compare_run,
)
from app.ordixa.tabular_profile import WORKFLOW_VERSION, load_table_profile
from app.ordixa.workbook import CHARTREUSE, FONT_NAME, GRAPHITE, MONEY_FORMAT

BEFORE = INVENTORY_DIR / "inventory-before.csv"
AFTER = INVENTORY_DIR / "inventory-after.csv"
INJECTION = ['=HYPERLINK("http://x")', "+cmd", "@SUM(A1)", "-2+3", "\tTab"]


def inventory_compare(out_dir, before=BEFORE, after=AFTER):
    """
    runs the fictional inventory comparison into a run directory
    :param out_dir: parent output directory
    :param before: earlier export
    :param after: later export
    :returns: tuple of run directory and manifest
    """
    profile = load_table_profile(COMPARE_PROFILE, "compare")
    return write_compare_run(compare_tables(profile, before, after), out_dir)


def read_csv(path):
    """
    reads a CSV artifact as dictionaries
    :param path: CSV path
    :returns: list of dictionaries
    """
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def test_clean_run_writes_cleaned_data_issues_and_manifest(tmp_path):
    """
    tests the clean run directory and its workflow-specific manifest
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    profile = load_table_profile(CLEAN_PROFILE, "clean")
    run_dir, manifest = write_clean_run(clean_table(profile, BEFORE), tmp_path)

    assert sorted(p.name for p in run_dir.iterdir()) == [
        "cleaned_data.csv",
        "cleaning_issues.csv",
        "manifest.json",
    ]
    rows = read_csv(run_dir / CLEANED_NAME)
    assert len(rows) == 120
    assert list(rows[0]) == [c.name for c in profile.columns]
    assert rows[0]["sku"] == "SKU-0001"
    assert [r["sku"] for r in rows if r["sku"].startswith("0")][:2] == [
        "000123",
        "000124",
    ]

    result = manifest["result"]
    assert result["workflow"] == "clean"
    assert result["workflow_version"] == WORKFLOW_VERSION
    assert result["profile"]["kind"] == "clean"
    assert result["profile"]["sha256"] == profile.sha256
    assert result["inputs"] == [
        {
            "role": "input",
            "file": "inventory-before.csv",
            "sha256": result["inputs"][0]["sha256"],
            "data_rows": 121,
        }
    ]
    assert result["row_counts"]["rows_written"] == 120
    assert result["issue_counts"] == {
        "Rejected": 0,
        "Removed duplicate": 1,
        "Note": 1,
    }
    assert result["status"] == "completed"
    assert all(output["deterministic"] for output in manifest["outputs"])
    assert all(passed for _, passed in verify_run(run_dir, [BEFORE]))


def test_compare_run_writes_report_csvs_and_manifest(tmp_path):
    """
    tests the compare run directory contents and manifest totals
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    run_dir, manifest = inventory_compare(tmp_path)
    assert sorted(p.name for p in run_dir.iterdir()) == [
        "added_records.csv",
        "comparison_issues.csv",
        "comparison_report.xlsx",
        "manifest.json",
        "modified_values.csv",
        "removed_records.csv",
    ]
    result = manifest["result"]
    assert result["workflow"] == "compare"
    assert result["status"] == "completed_with_rejections"
    assert [item["role"] for item in result["inputs"]] == ["before", "after"]
    assert result["ignored_columns"] == ["last_counted"]
    assert result["artifact_rows"] == {
        "added_records.csv": 6,
        "removed_records.csv": 4,
        "modified_values.csv": 20,
        "comparison_issues.csv": 8,
    }
    assert result["comparison"]["modified"] == 19
    assert [o["file"] for o in manifest["outputs"] if not o["deterministic"]] == [
        REPORT_NAME
    ]
    modified = read_csv(run_dir / MODIFIED_NAME)
    assert modified[0] == {
        "sku": "000125",
        "column": "unit_price",
        "before": "22.96",
        "after": "23.26",
        "difference": "0.30",
    }
    assert "Supplier Contact" not in (run_dir / ADDED_NAME).read_text(encoding="utf-8")
    assert "@" not in (run_dir / ADDED_NAME).read_text(encoding="utf-8")


def test_manifests_hold_no_row_values(tmp_path):
    """
    tests that manifests carry names, checksums, and counts but no cell values
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    run_dir, _ = inventory_compare(tmp_path)
    text = (run_dir / MANIFEST_NAME).read_text(encoding="utf-8")
    for value in ("SKU-0070", "Enamel", "23.26", "example", "Pocket Sketchbook"):
        assert value not in text


def test_reruns_are_deterministic_and_never_overwrite(tmp_path):
    """
    tests that two runs of the same inputs make separate directories with
    identical results, deterministic artifacts, and fingerprints
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    first_dir, first = inventory_compare(tmp_path)
    second_dir, second = inventory_compare(tmp_path)

    assert first_dir != second_dir and first_dir.exists() and second_dir.exists()
    assert first["result"] == second["result"]
    deterministic = {
        o["file"]: o["sha256"] for o in first["outputs"] if o["deterministic"]
    }
    assert deterministic == {
        o["file"]: o["sha256"] for o in second["outputs"] if o["deterministic"]
    }
    assert first["execution"]["run_id"] != second["execution"]["run_id"]


def test_fingerprint_reflects_the_logical_result(tmp_path):
    """
    tests that reordering rows keeps the fingerprint, and a real change
    alters it
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    rows = [["SKU-1", "Mug", "1.00", "1", ""], ["SKU-2", "Bowl", "2.00", "1", ""]]
    profile = table_profile()

    def fingerprint(name, after_rows):
        before = write_table(tmp_path / name / "before.csv", rows)
        after = write_table(tmp_path / name / "after.csv", after_rows)
        _, manifest = write_compare_run(
            compare_tables(profile, before, after), tmp_path / name / "out"
        )
        return manifest["result"]["result_fingerprint"]

    original = fingerprint("a", [rows[0], ["SKU-2", "Bowl", "2.50", "1", ""]])
    reordered = fingerprint("b", [["SKU-2", "Bowl", "$2.50", "1", ""], rows[0]])
    changed = fingerprint("c", [rows[0], ["SKU-2", "Bowl", "2.51", "1", ""]])
    assert original == reordered
    assert original != changed


def test_failed_writes_leave_no_run_directory(tmp_path, monkeypatch):
    """
    tests that a failure while writing publishes nothing
    :param tmp_path: pytest temporary directory
    :param monkeypatch: pytest monkeypatch fixture
    :returns: none
    """
    calls = []

    def failing_write_csv(path, columns, rows):
        calls.append(path)
        if len(calls) == 2:
            raise OSError("disk full")
        path.write_text("partial", encoding="utf-8")

    monkeypatch.setattr(tabular_outputs, "write_csv", failing_write_csv)
    with pytest.raises(OSError, match="disk full"):
        inventory_compare(tmp_path)
    assert list(tmp_path.iterdir()) == []


def test_verify_detects_tampering_and_changed_inputs(tmp_path):
    """
    tests verification of a compare run with inputs in two directories
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    week1 = tmp_path / "week1"
    week2 = tmp_path / "week2"
    week1.mkdir()
    week2.mkdir()
    before = shutil.copy(BEFORE, week1 / "inventory.csv")
    after = shutil.copy(AFTER, week2 / "inventory.csv")
    run_dir, _ = inventory_compare(tmp_path / "out", before, after)

    assert all(passed for _, passed in verify_run(run_dir, [week1, week2]))
    # one location only: the before file is there, the after file is not
    failed = [d for d, ok in verify_run(run_dir, [week1]) if not ok]
    assert failed == ["input inventory.csv is unchanged since the run"]

    with (run_dir / ADDED_NAME).open("a", encoding="utf-8") as handle:
        handle.write("SKU-9999,tampered\n")
    failed = [d for d, ok in verify_run(run_dir, [week1, week2]) if not ok]
    assert failed == [
        f"{ADDED_NAME} matches its recorded SHA-256 and size",
        f"{ADDED_NAME} has the recorded 6 data rows",
    ]


def test_formula_like_values_are_neutralized_everywhere(tmp_path):
    """
    tests that formula-like text is stored as inert text in the workbook and
    CSV artifacts, and that verify confirms no formulas
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    before = [[f"SKU-{i}", "Mug", "1.00", "1", ""] for i in range(len(INJECTION))]
    after = [[f"SKU-{i}", text, "1.00", "1", ""] for i, text in enumerate(INJECTION)]
    after.append(["SKU-NEW", "=1+1", "2.00", "1", ""])
    old = write_table(tmp_path / "before.csv", before)
    new = write_table(tmp_path / "after.csv", after)
    run_dir, _ = write_compare_run(
        compare_tables(table_profile(), old, new), tmp_path / "out"
    )

    workbook = load_workbook(run_dir / REPORT_NAME)
    cells = [cell for sheet in workbook for row in sheet.iter_rows() for cell in row]
    assert not [cell for cell in cells if cell.data_type == "f"]
    texts = {cell.value for cell in cells if isinstance(cell.value, str)}
    assert '\'=HYPERLINK("http://x")' in texts and "'=1+1" in texts
    assert "'@SUM(A1)" in texts and "'+cmd" in texts
    assert not [t for t in texts if t.startswith(("=", "+", "@"))]

    modified = read_csv(run_dir / MODIFIED_NAME)
    assert {row["after"] for row in modified} >= {"'+cmd", "'@SUM(A1)", "'-2+3"}
    assert read_csv(run_dir / ADDED_NAME)[0]["name"] == "'=1+1"
    assert all(passed for _, passed in verify_run(run_dir))


def test_change_report_layout_and_brand(tmp_path):
    """
    tests sheet order, headers, tab colors, typed cells, and print setup
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    run_dir, _ = inventory_compare(tmp_path)
    workbook = load_workbook(run_dir / REPORT_NAME)
    assert workbook.sheetnames == list(SHEET_NAMES)

    summary = workbook["Summary"]
    assert summary["A1"].value == "Ordixa"
    assert summary["A1"].font.color.rgb.endswith(CHARTREUSE)
    assert summary["A1"].fill.fgColor.rgb.endswith(GRAPHITE)
    labels = {
        summary.cell(row=r, column=1).value: summary.cell(row=r, column=2).value
        for r in range(1, summary.max_row + 1)
    }
    assert labels["Original records (rows in earlier export)"] == 121
    assert labels["New records (rows in later export)"] == 124
    assert (labels["Added"], labels["Removed"], labels["Modified"]) == (6, 4, 19)
    assert labels["Unchanged"] == 94
    assert labels["Ambiguous keys (conflicting rows)"] == 1
    assert labels["Invalid keys (a row was rejected)"] == 2

    modified = workbook["Modified Records"]
    assert [c.value for c in modified[1]] == [
        "sku",
        "Changed column",
        "Previous value",
        "New value",
        "Change",
    ]
    first = [c for c in modified[2]]
    assert first[0].value == "000125" and first[0].data_type == "s"
    assert first[2].number_format == MONEY_FORMAT
    assert first[4].value == 0.3 and first[4].number_format == MONEY_FORMAT

    added = workbook["Added Records"]
    skus = [added.cell(row=r, column=1).value for r in range(2, added.max_row + 1)]
    assert skus == ["000138", "000139", "BRM-016", "SKU-0091", "SKU-0092", "SKU-0093"]

    issues = workbook["Data Quality Issues"]
    assert issues.cell(row=2, column=1).value == "before"

    for sheet in workbook:
        assert sheet.sheet_properties.tabColor.rgb.endswith(TAB_COLORS[sheet.title])
        assert sheet.page_setup.fitToWidth == 1
        assert sheet.oddFooter.left.text == "Ordixa"
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    assert cell.font.name in (FONT_NAME, "Consolas")


def test_reconciliation_manifests_have_no_workflow_field(tmp_path):
    """
    tests that reconciliation manifests keep their original shape
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    from ordixa_helpers import EXPORTS_DIR, build_profile

    from app.ordixa.manifest import write_run_outputs
    from app.ordixa.pipeline import process

    _, manifest = write_run_outputs(
        process(build_profile(), EXPORTS_DIR / "clean"), tmp_path
    )
    assert "workflow" not in manifest["result"]
    assert manifest["manifest_version"] == 1
    assert json.loads(json.dumps(manifest)) == manifest
