"""
tests for workbook precision, identifier cells, artifact agreement, and run
directory safety
"""

import csv
import hashlib
import json
from datetime import datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from ordixa_golden import PROBLEMATIC_TOTALS, TEXT_ID_TOTALS
from ordixa_helpers import (
    EXPORTS_DIR,
    TEXT_ID_PROFILE_PATH,
    build_profile,
    write_exports,
)

from app.ordixa import adapters, manifest
from app.ordixa.cli import EXIT_INPUT, EXIT_OUTPUT, EXIT_UNEXPECTED, main
from app.ordixa.manifest import (
    ISSUES_NAME,
    LEDGER_NAME,
    MANIFEST_NAME,
    WORKBOOK_NAME,
    write_run_outputs,
)
from app.ordixa.pipeline import process
from app.ordixa.profile import load_profile
from app.ordixa.workbook import MAX_CELL_TEXT, exact_in_excel

D = Decimal
TEXT_MODES = {"customer_id": "text", "order_id": "text", "payment_id": "text"}


def sheet_rows(sheet):
    """
    reads a table sheet into dictionaries keyed by header
    :param sheet: worksheet
    :returns: list of dictionaries
    """
    rows = list(sheet.iter_rows(values_only=True))
    return [dict(zip(rows[0], row)) for row in rows[1:]]


def summary_cells(sheet):
    """
    maps summary labels to their value cells
    :param sheet: summary worksheet
    :returns: dictionary of label to cell
    """
    return {
        row[0].value: row[1]
        for row in sheet.iter_rows(min_col=1, max_col=2)
        if row[0].value is not None
    }


def csv_rows(path):
    """
    reads a CSV artifact
    :param path: CSV path
    :returns: list of dictionaries
    """
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def order(ref, total="10.00", customer="7000101", when="2026-08-01"):
    """
    builds an order row
    :param ref: order name
    :param total: total text
    :param customer: customer id text
    :param when: created at text
    :returns: list of cells
    """
    return [ref, when, customer, "USD", total]


@pytest.mark.parametrize(
    ("value", "exact"),
    [
        (D("9999999999.99"), True),
        (D("9999999999999.99"), True),
        (D("99999999999999.99"), False),
        (D("0.01"), True),
        (10**15 - 1, True),
        (10**15, False),
    ],
)
def test_excel_precision_boundary(value, exact):
    """
    15 significant digits is the largest exact Excel number
    :param value: number to check
    :param exact: expected result
    :returns: none
    """
    assert exact_in_excel(value) is exact


def test_totals_beyond_excel_precision_are_written_as_exact_text(tmp_path):
    """
    hand-computed: 1,001 orders of 9,999,999,999.99 total
    9,999,999,999.99 x 1,001 = 10,009,999,999,989.99 (16 significant digits);
    with no payments the outstanding balance is the same; both are stored as
    exact text instead of being rounded by Excel
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    input_dir = write_exports(
        tmp_path / "in",
        orders=[order(f"#{n}", "9,999,999,999.99") for n in range(1, 1002)],
        payments=[],
    )

    run_dir, run_manifest = write_run_outputs(
        process(build_profile(), input_dir), tmp_path / "out"
    )
    workbook = load_workbook(run_dir / WORKBOOK_NAME)
    cells = summary_cells(workbook["Summary"])

    assert cells["Order value"].value == "10009999999989.99"
    assert cells["Order value"].data_type == "s"
    assert cells["Outstanding balances"].value == "10009999999989.99"
    assert cells["Amounts stored as exact text (over 15 digits)"].value == 2
    assert run_manifest["result"]["totals"]["order_value"] == "10009999999989.99"
    ledger = sheet_rows(workbook["Full Ledger"])
    assert {row["Order total"] for row in ledger} == {9999999999.99}


def test_ordinary_workbooks_have_no_precision_note(tmp_path):
    """
    the precision line only appears when it is needed
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    run_dir, _ = write_run_outputs(
        process(build_profile(), EXPORTS_DIR / "clean"), tmp_path
    )
    cells = summary_cells(load_workbook(run_dir / WORKBOOK_NAME)["Summary"])

    assert "Amounts stored as exact text (over 15 digits)" not in cells
    assert isinstance(cells["Order value"].value, float)


def test_identifiers_are_text_cells_in_every_artifact(tmp_path):
    """
    a 19-digit numeric order id and text ids with leading zeros survive the
    workbook and CSV exactly
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    long_id = "1234567890123456789"
    numeric_dir = write_exports(
        tmp_path / "numeric",
        orders=[order(f"#{long_id}")],
        payments=[],
    )
    text_dir = write_exports(
        tmp_path / "text",
        orders=[order("#00012", customer="0007000101")],
        payments=[],
    )
    text_profile = build_profile(lambda d: d.update(identifiers=dict(TEXT_MODES)))

    numeric_run, _ = write_run_outputs(
        process(build_profile(), numeric_dir), tmp_path / "out"
    )
    text_run, _ = write_run_outputs(process(text_profile, text_dir), tmp_path / "out")

    numeric_sheet = load_workbook(numeric_run / WORKBOOK_NAME)["Full Ledger"]
    text_sheet = load_workbook(text_run / WORKBOOK_NAME)["Full Ledger"]
    assert sheet_rows(numeric_sheet)[0]["Order ID"] == long_id
    assert sheet_rows(text_sheet)[0]["Order ID"] == "00012"
    assert sheet_rows(text_sheet)[0]["Customer ID"] == "0007000101"
    assert csv_rows(numeric_run / LEDGER_NAME)[0]["order_id"] == long_id
    assert csv_rows(text_run / LEDGER_NAME)[0]["order_id"] == "00012"
    assert csv_rows(text_run / LEDGER_NAME)[0]["customer_id"] == "0007000101"


def test_long_and_unusual_text_is_kept_safely(tmp_path):
    """
    text over Excel's cell limit is truncated with a marker, while emoji and
    right-to-left text survive and control characters are removed
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    long_name = "N" * 40_000
    input_dir = write_exports(
        tmp_path / "in",
        orders=[order("#1"), order("#2", customer="7000102")],
        payments=[],
        customers=[
            ["7000101", long_name, "", "a@example.com", "555-010-0101"],
            ["7000102", "Zoë 🌿 שלום", "Bell\x07", "b@example.com", "555-010-0102"],
        ],
    )

    run_dir, _ = write_run_outputs(
        process(build_profile(), input_dir), tmp_path / "out"
    )
    ledger = sheet_rows(load_workbook(run_dir / WORKBOOK_NAME)["Full Ledger"])

    assert len(ledger[0]["Customer"]) == MAX_CELL_TEXT
    assert ledger[0]["Customer"].endswith("[truncated]")
    assert ledger[1]["Customer"] == "Zoë 🌿 שלום Bell"
    assert len(csv_rows(run_dir / LEDGER_NAME)[0]["customer_name"]) == 40_000


def test_formula_like_identifiers_cannot_become_formulas(tmp_path):
    """
    a rejected identifier echoed into the Rejected Rows sheet stays inert
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    input_dir = write_exports(
        tmp_path / "in",
        orders=[order("#1")],
        payments=[
            ['=HYPERLINK("x")', "#1", "08/02/2026", "Payment", "Paid", "$1.00", "USD"]
        ],
    )
    profile = build_profile(lambda d: d.update(identifiers=dict(TEXT_MODES)))
    profile.datasets["payments"].id_prefixes.clear()

    run_dir, _ = write_run_outputs(process(profile, input_dir), tmp_path / "out")
    sheet = load_workbook(run_dir / WORKBOOK_NAME)["Rejected Rows"]
    cell = next(c for c in sheet["D"] if c.value and "HYPERLINK" in c.value)

    assert cell.data_type == "s"
    assert cell.value.startswith("'=")
    assert csv_rows(run_dir / ISSUES_NAME)[0]["record_id"].startswith("'=")


def test_offset_timestamps_keep_their_calendar_date(tmp_path):
    """
    23:30 at -0400 on 31 August is 1 September in UTC; the export's own
    calendar date is what the workbook shows
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    input_dir = write_exports(
        tmp_path / "in",
        orders=[order("#1", when="2026-08-31 23:30:00 -0400")],
        payments=[],
    )

    run_dir, _ = write_run_outputs(
        process(build_profile(), input_dir), tmp_path / "out"
    )
    row = sheet_rows(load_workbook(run_dir / WORKBOOK_NAME)["Full Ledger"])[0]

    assert row["Order date"] == datetime(2026, 8, 31)


@pytest.mark.parametrize(
    ("profile_path", "input_name", "totals"),
    [
        (None, "problematic", PROBLEMATIC_TOTALS),
        (TEXT_ID_PROFILE_PATH, "text-ids", TEXT_ID_TOTALS),
    ],
)
def test_counts_and_totals_agree_across_every_artifact(
    tmp_path, profile_path, input_name, totals
):
    """
    workbook sheets, CSV artifacts, and the manifest describe the same run, and
    the totals equal the hand-computed values
    :param tmp_path: pytest temporary directory
    :param profile_path: profile to load, or none for Harbor Lane
    :param input_name: fixture directory
    :param totals: hand-computed totals
    :returns: none
    """
    profile = load_profile(profile_path) if profile_path else build_profile()
    run_dir, run_manifest = write_run_outputs(
        process(profile, EXPORTS_DIR / input_name), tmp_path
    )
    workbook = load_workbook(run_dir / WORKBOOK_NAME)
    result = run_manifest["result"]
    ledger_csv = csv_rows(run_dir / LEDGER_NAME)
    issues_csv = csv_rows(run_dir / ISSUES_NAME)
    ledger_sheet = sheet_rows(workbook["Full Ledger"])
    rejected_sheet = sheet_rows(workbook["Rejected Rows"])

    assert len(ledger_csv) == len(ledger_sheet) == result["artifact_rows"][LEDGER_NAME]
    assert len(ledger_csv) == result["row_counts"]["orders"]["accepted"]
    assert (
        len(issues_csv) == len(rejected_sheet) == result["artifact_rows"][ISSUES_NAME]
    )
    assert len(issues_csv) == sum(result["issue_counts"].values())
    rejected = sum(1 for row in rejected_sheet if row["Outcome"] == "Rejected")
    assert rejected == result["issue_counts"]["error"]

    cells = summary_cells(workbook["Summary"])
    labels = {
        "order_value": "Order value",
        "net_received": "Net received",
        "outstanding": "Outstanding balances",
        "outstanding_balance": "Outstanding balances",
        "overpaid": "Overpayments",
        "overpayments": "Overpayments",
    }
    for key, expected in totals.items():
        if key not in labels:
            continue
        manifest_key = {
            "outstanding": "outstanding_balance",
            "overpaid": "overpayments",
        }
        assert D(result["totals"][manifest_key.get(key, key)]) == expected, key
        assert D(repr(cells[labels[key]].value)).quantize(D("0.01")) == expected, key
    net_from_csv = sum((D(row["net_paid"]) for row in ledger_csv), D("0.00"))
    assert net_from_csv == D(result["totals"]["net_received"])


def test_manifest_records_identifier_modes_without_row_values(tmp_path):
    """
    the manifest says how identifiers were read but holds no identifiers
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    run_dir, _ = write_run_outputs(
        process(load_profile(TEXT_ID_PROFILE_PATH), EXPORTS_DIR / "text-ids"),
        tmp_path,
    )
    text = (run_dir / MANIFEST_NAME).read_text("utf-8")

    assert json.loads(text)["result"]["profile"]["identifiers"] == TEXT_MODES
    for fragment in ("00012345", "ORD-2026", "PAY_8F3K2", "Morgan", "@example"):
        assert fragment not in text


def test_input_checksums_cover_the_exact_bytes_read(tmp_path):
    """
    a CSV with a byte-order mark and CRLF line endings is fingerprinted as
    stored on disk
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    input_dir = tmp_path / "in"
    write_exports(input_dir, orders=[order("#1")], payments=[])
    orders = input_dir / "orders.csv"
    raw = b"\xef\xbb\xbf" + orders.read_bytes().replace(b"\n", b"\r\n")
    orders.write_bytes(raw)

    _, run_manifest = write_run_outputs(
        process(build_profile(), input_dir), tmp_path / "out"
    )
    recorded = {
        item["file"]: item["sha256"] for item in run_manifest["result"]["inputs"]
    }

    assert recorded["orders.csv"] == hashlib.sha256(raw).hexdigest()


def test_oversized_inputs_are_refused(tmp_path, monkeypatch, capsys):
    """
    the input size guard stops the run with an input error
    :param tmp_path: pytest temporary directory
    :param monkeypatch: pytest monkeypatch fixture
    :param capsys: pytest capture fixture
    :returns: none
    """
    monkeypatch.setattr(adapters, "MAX_INPUT_BYTES", 10)
    input_dir = write_exports(tmp_path / "in", orders=[order("#1")], payments=[])

    code = main(
        [
            "run",
            "--profile",
            str(EXPORTS_DIR.parent / "profiles" / "harbor-lane.toml"),
            "--input",
            str(input_dir),
            "--out",
            str(tmp_path / "out"),
        ]
    )

    assert code == EXIT_INPUT
    assert "input limit" in capsys.readouterr().err
    assert not (tmp_path / "out").exists()


@pytest.mark.parametrize(
    ("error", "exit_code"),
    [
        (OSError(28, "No space left on device"), EXIT_OUTPUT),
        (RuntimeError("boom"), EXIT_UNEXPECTED),
    ],
)
def test_failed_writes_leave_no_run_directory(tmp_path, monkeypatch, error, exit_code):
    """
    a failure after some artifacts were written removes the staging directory,
    so nothing that looks like a completed run is left behind
    :param tmp_path: pytest temporary directory
    :param monkeypatch: pytest monkeypatch fixture
    :param error: exception raised while writing
    :param exit_code: expected CLI exit code
    :returns: none
    """
    original = manifest.write_csv
    calls = []

    def failing_write_csv(path, columns, rows):
        original(path, columns, rows)
        calls.append(path)
        if len(calls) == 2:
            raise error

    monkeypatch.setattr(manifest, "write_csv", failing_write_csv)
    out_dir = tmp_path / "out"

    code = main(
        [
            "run",
            "--profile",
            str(EXPORTS_DIR.parent / "profiles" / "harbor-lane.toml"),
            "--input",
            str(EXPORTS_DIR / "clean"),
            "--out",
            str(out_dir),
        ]
    )

    assert code == exit_code
    assert len(calls) == 2
    assert list(out_dir.iterdir()) == []


def test_successful_runs_leave_no_staging_directory(tmp_path):
    """
    only the finished run-<id> directory remains, with its manifest
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    run_dir, run_manifest = write_run_outputs(
        process(build_profile(), EXPORTS_DIR / "clean"), tmp_path
    )

    assert [path.name for path in tmp_path.iterdir()] == [run_dir.name]
    assert run_manifest["execution"]["output_directory"] == run_dir.name
    assert sorted(path.name for path in run_dir.iterdir()) == sorted(
        [WORKBOOK_NAME, LEDGER_NAME, ISSUES_NAME, MANIFEST_NAME]
    )
