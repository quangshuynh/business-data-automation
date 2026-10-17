"""
tests that runs written before the 2.0.0 release still verify, and that the
2.0.0 engine reproduces their deterministic outputs byte for byte

tests/fixtures/runs-0.1.0 holds unmodified run directories produced by the
pre-release engine (application version 0.1.0, commit 00b86d3) from the
fictional example exports; they are read here and never rewritten
"""

import json

import pytest
from ordixa_helpers import EXPORTS_DIR, ROOT

from app.ordixa.manifest import MANIFEST_NAME, verify_run, write_run_outputs
from app.ordixa.pipeline import process
from app.ordixa.profile import load_profile
from app.ordixa.tabular import clean_table, compare_tables
from app.ordixa.tabular_outputs import write_clean_run, write_compare_run
from app.ordixa.tabular_profile import load_table_profile

FIXTURES = ROOT / "tests" / "fixtures" / "runs-0.1.0"
PROFILES = ROOT / "examples" / "ordixa" / "profiles"
INVENTORY = EXPORTS_DIR / "inventory"


def load_manifest(run_dir):
    """
    reads a run manifest
    :param run_dir: run directory
    :returns: manifest dictionary
    """
    return json.loads((run_dir / MANIFEST_NAME).read_text(encoding="utf-8"))


def reproduce(name, out_dir):
    """
    reruns one fixture's workflow with the current engine
    :param name: fixture directory name
    :param out_dir: parent directory for the new run
    :returns: tuple of new run directory and manifest
    """
    if name == "reconcile-harbor-lane":
        profile = load_profile(PROFILES / "harbor-lane.toml")
        return write_run_outputs(process(profile, EXPORTS_DIR / "problematic"), out_dir)
    if name == "reconcile-juniper-row":
        profile = load_profile(PROFILES / "juniper-row.toml")
        return write_run_outputs(process(profile, EXPORTS_DIR / "text-ids"), out_dir)
    if name == "clean-inventory":
        profile = load_table_profile(PROFILES / "inventory-clean.toml")
        result = clean_table(profile, INVENTORY / "inventory-before.csv")
        return write_clean_run(result, out_dir)
    profile = load_table_profile(PROFILES / "inventory-compare.toml")
    result = compare_tables(
        profile, INVENTORY / "inventory-before.csv", INVENTORY / "inventory-after.csv"
    )
    return write_compare_run(result, out_dir)


INPUTS = {
    "reconcile-harbor-lane": [EXPORTS_DIR / "problematic"],
    "reconcile-juniper-row": [EXPORTS_DIR / "text-ids"],
    "clean-inventory": [INVENTORY / "inventory-before.csv"],
    "compare-inventory": [INVENTORY],
}


@pytest.mark.parametrize("name", sorted(INPUTS))
def test_pre_release_runs_still_verify(name):
    """
    tests that verify accepts manifest version 1 runs from 0.1.0 unchanged
    :param name: fixture directory name
    :returns: none
    """
    run_dir = FIXTURES / name
    manifest = load_manifest(run_dir)
    assert manifest["manifest_version"] == 1
    assert manifest["result"]["code_version"]["ordixa"] == "0.1.0"
    checks = verify_run(run_dir, INPUTS[name])
    assert checks and all(passed for _, passed in checks), checks


@pytest.mark.parametrize("name", sorted(INPUTS))
def test_current_engine_reproduces_pre_release_outputs(name, tmp_path):
    """
    tests that 2.0.0 keeps result fingerprints and deterministic artifacts
    :param name: fixture directory name
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    old = load_manifest(FIXTURES / name)
    _, new = reproduce(name, tmp_path)

    assert new["manifest_version"] == old["manifest_version"]
    assert new["result"]["result_fingerprint"] == old["result"]["result_fingerprint"]
    assert new["result"]["status"] == old["result"]["status"]
    assert new["result"]["artifact_rows"] == old["result"]["artifact_rows"]
    assert new["result"]["inputs"] == old["result"]["inputs"]
    deterministic = {
        output["file"]: output["sha256"]
        for output in old["outputs"]
        if output["deterministic"]
    }
    assert deterministic
    assert {
        output["file"]: output["sha256"]
        for output in new["outputs"]
        if output["deterministic"]
    } == deterministic
    assert [o["file"] for o in new["outputs"]] == [o["file"] for o in old["outputs"]]


@pytest.mark.parametrize("name", sorted(INPUTS))
def test_manifest_fields_from_0_1_0_are_all_still_written(name, tmp_path):
    """
    tests that 2.0.0 manifests only add fields, never drop or rename them
    :param name: fixture directory name
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    old = load_manifest(FIXTURES / name)
    _, new = reproduce(name, tmp_path)
    assert set(old) <= set(new)
    for section in ("result", "execution"):
        assert set(old[section]) <= set(new[section]), section


# Summary rows that describe the run rather than its result
RUN_SPECIFIC_LABELS = {"Run ID", "Generated (UTC)", "Application version"}


def sheet_values(sheet):
    """
    reads every cell value of a sheet, leaving out run-specific metadata
    :param sheet: worksheet
    :returns: list of row tuples
    """
    rows = []
    for row in sheet.iter_rows(values_only=True):
        if row and row[0] in RUN_SPECIFIC_LABELS:
            continue
        if row and isinstance(row[0], str) and " generated " in row[0]:
            continue
        rows.append(row)
    return rows


@pytest.mark.parametrize(
    "name",
    ["reconcile-harbor-lane", "reconcile-juniper-row", "compare-inventory"],
)
def test_workbook_values_match_pre_release_runs(name, tmp_path):
    """
    tests that every workbook figure and record is unchanged since 0.1.0

    the reconciliation workbook is compared sheet by sheet in full; the change
    report gained a Comparison result section on its Summary sheet in 2.0.0,
    so only its record sheets are compared cell for cell
    :param name: fixture directory name
    :param tmp_path: pytest temporary directory
    :returns: none
    """
    from openpyxl import load_workbook

    old_dir = FIXTURES / name
    new_dir, manifest = reproduce(name, tmp_path)
    (workbook_name,) = [
        o["file"] for o in manifest["outputs"] if o["file"].endswith(".xlsx")
    ]
    old = load_workbook(old_dir / workbook_name)
    new = load_workbook(new_dir / workbook_name)
    assert old.sheetnames == new.sheetnames
    for title in old.sheetnames:
        if name == "compare-inventory" and title == "Summary":
            continue
        assert sheet_values(new[title]) == sheet_values(old[title]), title
    old.close()
    new.close()
